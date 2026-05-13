#!/usr/bin/env python3
"""Build a manual review workbook for excluding corporate/external groups."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_PIPELINE_ROOT = Path("/home/thiago/1_Supervised_BERTopic/paper_6topic_discourse_pipeline")
DEFAULT_OUTPUT_DIR = (
    DEFAULT_PIPELINE_ROOT
    / "outputs"
    / "corporate_external_group_exclusion_review"
)

APPENDIX_TABLES = (
    Path("outputs")
    / "paper_tables"
    / "appendix_method_tables"
    / "paper_appendix_method_tables.xlsx"
)
ALL_GROUPS = (
    Path("outputs")
    / "corporate_focus_review_with_overrides"
    / "corporate_focus_all_groups_optional.csv"
)
RELATION_TABLES = (
    Path("outputs")
    / "paper_tables"
    / "corporate_focus_relative_longitudinal_series"
    / "topic_tables"
)

REVIEW_DECISIONS = ["keep", "exclude"]
EXCLUSION_REASONS = [
    "not_sustainability",
    "generic_business",
    "generic_financial",
    "generic_competition",
    "wrong_sector_or_scope",
    "duplicate_or_absorbed",
    "other",
]
EDITABLE_COLUMNS = ["review_decision", "exclusion_reason_category", "review_notes"]
INFO_COLUMNS = [
    "topic_id",
    "macro_topic",
    "source",
    "subgroup",
    "final_merge_group_id",
    "topic_label",
    "overall_summary",
    "topic_size",
    "active_year_min",
    "active_year_max",
    "current_selection_bucket",
    "best_matched_corporate_group_id",
    "academic_counterpart_count",
    "media_counterpart_count",
    "counterpart_labels",
]
REVIEW_COLUMNS = [*EDITABLE_COLUMNS, *INFO_COLUMNS]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pipeline-root", type=Path, default=DEFAULT_PIPELINE_ROOT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args()


def clean_text(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return " ".join(str(value).split())


def make_topic_id(subgroup: object, group_id: object) -> str:
    return f"{clean_text(subgroup)}::{clean_text(group_id)}"


def split_pipe(value: object) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    return [part.strip() for part in text.split(" | ") if part.strip()]


def join_labeled_counterparts(row: pd.Series) -> str:
    parts: list[str] = []
    for source in ["academic", "media"]:
        labels = split_pipe(row.get(f"{source}_labels", ""))
        parts.extend(f"{source}: {label}" for label in labels)
    return " | ".join(parts)


def load_all_groups(pipeline_root: Path) -> pd.DataFrame:
    path = pipeline_root / ALL_GROUPS
    groups = pd.read_csv(path).fillna("")
    groups["topic_id"] = [make_topic_id(row.subgroup, row.final_merge_group_id) for row in groups.itertuples(index=False)]
    groups["topic_label"] = groups["topic_label_refined"].map(clean_text)
    fallback = groups["topic_name_original"].map(clean_text)
    groups.loc[groups["topic_label"] == "", "topic_label"] = fallback
    groups["current_selection_bucket"] = groups["selection_bucket"].map(clean_text)
    return groups


def load_corporate_aggregates(pipeline_root: Path) -> pd.DataFrame:
    path = pipeline_root / APPENDIX_TABLES
    aggregates = pd.read_excel(path, sheet_name="Corporate with aggregates").fillna("")
    aggregates["topic_id"] = [
        make_topic_id(row.subgroup_corporate, row.final_merge_group_id_corporate)
        for row in aggregates.itertuples(index=False)
    ]
    aggregates["academic_counterpart_count"] = pd.to_numeric(
        aggregates["academic_topic_count"], errors="coerce"
    ).fillna(0).astype(int)
    aggregates["media_counterpart_count"] = pd.to_numeric(
        aggregates["media_topic_count"], errors="coerce"
    ).fillna(0).astype(int)
    aggregates["counterpart_labels"] = aggregates.apply(join_labeled_counterparts, axis=1)
    return aggregates[
        [
            "topic_id",
            "academic_counterpart_count",
            "media_counterpart_count",
            "counterpart_labels",
        ]
    ]


def load_pair_context(pipeline_root: Path) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for path in sorted((pipeline_root / RELATION_TABLES).glob("T*_aligned_aggregate_topics.csv")):
        frame = pd.read_csv(path).fillna("")
        if frame.empty:
            continue
        frame["source_file"] = str(path.relative_to(pipeline_root))
        frame["corporate_topic_id"] = [
            make_topic_id(row.subgroup_corporate, row.final_merge_group_id_corporate)
            for row in frame.itertuples(index=False)
        ]
        frame["external_topic_id"] = [
            make_topic_id(row.subgroup_noncorporate, row.final_merge_group_id_noncorporate)
            for row in frame.itertuples(index=False)
        ]
        frames.append(frame)
    if not frames:
        return pd.DataFrame()
    pairs = pd.concat(frames, ignore_index=True)
    columns = [
        "macro_topic",
        "source_noncorporate",
        "external_topic_id",
        "subgroup_noncorporate",
        "final_merge_group_id_noncorporate",
        "topic_label_refined_noncorporate",
        "external_unique_document_count",
        "corporate_topic_id",
        "subgroup_corporate",
        "final_merge_group_id_corporate",
        "topic_label_refined_corporate",
        "corporate_unique_document_count",
        "best_cosine_similarity",
        "direct_pair_count",
        "best_pair_id",
        "inclusion_reason",
        "source_file",
    ]
    return pairs[[column for column in columns if column in pairs.columns]].drop_duplicates()


def pair_lookup(pair_context: pd.DataFrame) -> dict[str, dict[str, Any]]:
    if pair_context.empty:
        return {}
    lookup: dict[str, dict[str, Any]] = {}
    for external_topic_id, group in pair_context.groupby("external_topic_id", sort=False):
        labels = [
            clean_text(label)
            for label in group["topic_label_refined_corporate"].dropna().astype(str)
            if clean_text(label)
        ]
        ids = [
            clean_text(topic_id)
            for topic_id in group["corporate_topic_id"].dropna().astype(str)
            if clean_text(topic_id)
        ]
        lookup[str(external_topic_id)] = {
            "paired_corporate_count": int(group["corporate_topic_id"].nunique()),
            "paired_corporate_topic_ids": " | ".join(dict.fromkeys(ids)),
            "paired_corporate_labels": " | ".join(dict.fromkeys(labels)),
        }
    return lookup


def empty_review_columns(frame: pd.DataFrame) -> pd.DataFrame:
    updated = frame.copy()
    updated.insert(0, "review_decision", "")
    updated.insert(1, "exclusion_reason_category", "")
    updated.insert(2, "review_notes", "")
    return updated


def build_corporate_review(groups: pd.DataFrame, aggregates: pd.DataFrame) -> pd.DataFrame:
    corporate = groups.loc[groups["source"] == "corporate"].copy()
    corporate = corporate.merge(aggregates, on="topic_id", how="left")
    for column in ["academic_counterpart_count", "media_counterpart_count"]:
        corporate[column] = pd.to_numeric(corporate[column], errors="coerce").fillna(0).astype(int)
    corporate["counterpart_labels"] = corporate["counterpart_labels"].fillna("").map(clean_text)
    corporate["best_matched_corporate_group_id"] = ""
    review = corporate[INFO_COLUMNS].sort_values(["macro_topic", "topic_label"], kind="mergesort")
    return empty_review_columns(review)


def build_external_review(groups: pd.DataFrame, pairs: pd.DataFrame) -> pd.DataFrame:
    external = groups.loc[groups["source"].isin(["academic", "media"])].copy()
    external["academic_counterpart_count"] = 0
    external["media_counterpart_count"] = 0
    lookup = pair_lookup(pairs)
    external["paired_corporate_count"] = [lookup.get(topic_id, {}).get("paired_corporate_count", 0) for topic_id in external["topic_id"]]
    external["paired_corporate_topic_ids"] = [
        lookup.get(topic_id, {}).get("paired_corporate_topic_ids", "") for topic_id in external["topic_id"]
    ]
    external["paired_corporate_labels"] = [
        lookup.get(topic_id, {}).get("paired_corporate_labels", "") for topic_id in external["topic_id"]
    ]
    external["counterpart_labels"] = external["paired_corporate_labels"].map(clean_text)
    review = external[
        [
            *INFO_COLUMNS,
            "paired_corporate_count",
            "paired_corporate_topic_ids",
            "paired_corporate_labels",
            "noncorporate_inclusion_reason",
            "best_dyad",
        ]
    ].sort_values(["macro_topic", "source", "topic_label"], kind="mergesort")
    return empty_review_columns(review)


def build_readme() -> pd.DataFrame:
    rows = [
        {
            "section": "Purpose",
            "detail": "Manual review workbook for marking corporate, media, and academic groups to exclude from the paper/app because they are not directly and explicitly about sustainability.",
        },
        {
            "section": "How to mark",
            "detail": "Use review_decision = exclude. Leave review_decision blank, or set it to keep, to retain a group.",
        },
        {
            "section": "Reason categories",
            "detail": ", ".join(EXCLUSION_REASONS),
        },
        {
            "section": "Decision scope",
            "detail": "An exclude decision means remove from paper + app in a later materialization step. This workbook only records decisions.",
        },
        {
            "section": "Origin",
            "detail": "The Appendix II document corresponds to Appendix Table A3 generated from paper_appendix_method_tables.xlsx / Corporate with aggregates.",
        },
        {
            "section": "Editable sheets",
            "detail": "Edit only Corporate groups and External groups. Pair context is read-only context, and Decision export is a normalized empty template.",
        },
    ]
    return pd.DataFrame(rows)


def build_decision_export() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            "review_decision",
            "exclusion_reason_category",
            "review_notes",
            "topic_id",
            "macro_topic",
            "source",
            "subgroup",
            "final_merge_group_id",
            "topic_label",
        ]
    )


def autosize_and_style(path: Path, review_sheet_names: list[str]) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="B7D7F0")
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(bold=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = review_fill if cell.value in EDITABLE_COLUMNS else header_fill
            cell.font = header_font
            cell.alignment = body_alignment
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = body_alignment
        for column_cells in sheet.columns:
            header = clean_text(column_cells[0].value)
            max_len = max(len(clean_text(cell.value)) for cell in column_cells[:200])
            width = min(max(max_len + 2, len(header) + 2, 12), 70)
            sheet.column_dimensions[column_cells[0].column_letter].width = width

    decision_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(REVIEW_DECISIONS)}"',
        allow_blank=True,
    )
    reason_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(EXCLUSION_REASONS)}"',
        allow_blank=True,
    )
    for sheet_name in review_sheet_names:
        sheet = workbook[sheet_name]
        sheet.add_data_validation(decision_validation)
        sheet.add_data_validation(reason_validation)
        max_row = max(sheet.max_row, 2)
        decision_validation.add(f"A2:A{max_row}")
        reason_validation.add(f"B2:B{max_row}")

    workbook.save(path)


def write_workbook(
    output_path: Path,
    corporate_review: pd.DataFrame,
    external_review: pd.DataFrame,
    pair_context_frame: pd.DataFrame,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        build_readme().to_excel(writer, sheet_name="README", index=False)
        corporate_review.to_excel(writer, sheet_name="Corporate groups", index=False)
        external_review.to_excel(writer, sheet_name="External groups", index=False)
        pair_context_frame.to_excel(writer, sheet_name="Pair context", index=False)
        build_decision_export().to_excel(writer, sheet_name="Decision export", index=False)
    autosize_and_style(output_path, ["Corporate groups", "External groups"])


def write_manifest(
    output_dir: Path,
    pipeline_root: Path,
    corporate_review: pd.DataFrame,
    external_review: pd.DataFrame,
    pair_context_frame: pd.DataFrame,
    workbook_path: Path,
) -> None:
    manifest = {
        "workbook": str(workbook_path),
        "pipeline_root": str(pipeline_root),
        "input_files": {
            "all_groups_optional": str(pipeline_root / ALL_GROUPS),
            "appendix_method_tables": str(pipeline_root / APPENDIX_TABLES),
            "relation_topic_tables": str(pipeline_root / RELATION_TABLES),
        },
        "row_counts": {
            "corporate_groups": int(len(corporate_review)),
            "external_groups": int(len(external_review)),
            "pair_context": int(len(pair_context_frame)),
            "corporate_without_counterparts": int(
                (
                    (corporate_review["academic_counterpart_count"] == 0)
                    & (corporate_review["media_counterpart_count"] == 0)
                ).sum()
            ),
        },
        "editable_columns": EDITABLE_COLUMNS,
        "review_decision_values": REVIEW_DECISIONS,
        "exclusion_reason_values": EXCLUSION_REASONS,
    }
    (output_dir / "corporate_external_group_exclusion_review_manifest.json").write_text(
        json.dumps(manifest, indent=2),
        encoding="utf-8",
    )


def main() -> None:
    args = parse_args()
    groups = load_all_groups(args.pipeline_root)
    aggregates = load_corporate_aggregates(args.pipeline_root)
    pairs = load_pair_context(args.pipeline_root)

    corporate_review = build_corporate_review(groups, aggregates)
    external_review = build_external_review(groups, pairs)
    output_path = args.output_dir / "corporate_external_group_exclusion_review.xlsx"
    write_workbook(output_path, corporate_review, external_review, pairs)
    write_manifest(args.output_dir, args.pipeline_root, corporate_review, external_review, pairs, output_path)
    print(
        json.dumps(
            {
                "workbook": str(output_path),
                "corporate_groups": int(len(corporate_review)),
                "external_groups": int(len(external_review)),
                "pair_context": int(len(pairs)),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
