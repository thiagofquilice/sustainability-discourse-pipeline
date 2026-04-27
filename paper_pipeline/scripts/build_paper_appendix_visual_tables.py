#!/usr/bin/env python3
"""Build reader-oriented appendix tables from the full method workbook."""

from __future__ import annotations

from pathlib import Path
from textwrap import shorten

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PIPELINE_ROOT = Path("paper_pipeline")
APPENDIX_DIR = PIPELINE_ROOT / "outputs" / "paper_tables" / "appendix_method_tables"
FULL_XLSX = APPENDIX_DIR / "paper_appendix_method_tables.xlsx"
OUTPUT_XLSX = APPENDIX_DIR / "paper_appendix_visual_tables.xlsx"
OUTPUT_MD = APPENDIX_DIR / "paper_appendix_visual_tables.md"

MACRO_ORDER = ["T1", "T2", "T3", "T4", "T5", "T6"]

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FILL = PatternFill("solid", fgColor="B7D7F0")
LIGHT_FILL = PatternFill("solid", fgColor="F5F9FC")
WHITE_FONT = Font(name="Times New Roman", color="FFFFFF", bold=True, size=13)
HEADER_FONT = Font(name="Times New Roman", bold=True, size=11)
BODY_FONT = Font(name="Times New Roman", size=10)
SMALL_FONT = Font(name="Times New Roman", size=9)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)


def clean_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return " ".join(str(value).split())


def paragraph(value: object, limit: int = 420) -> str:
    text = clean_text(value)
    if not text:
        return ""
    return shorten(text, width=limit, placeholder="...")


def split_pipe(value: object) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    return [part.strip() for part in text.split(" | ") if part.strip()]


def bullet_lines(items: list[str], empty_text: str = "No retained counterpart.") -> str:
    if not items:
        return empty_text
    return "\n".join(f"- {item}" for item in items)


def counterpart_cell(row: pd.Series, source: str) -> str:
    topic_count = int(pd.to_numeric(row.get(f"{source}_topic_count"), errors="coerce") or 0)
    doc_count = int(pd.to_numeric(row.get(f"{source}_document_count"), errors="coerce") or 0)
    labels = split_pipe(row.get(f"{source}_labels"))
    if topic_count <= 0 or not labels:
        return "No retained counterpart."
    header = f"{topic_count} topics; {doc_count} unique documents"
    return f"{header}\n{bullet_lines(labels)}"


def counterpart_cell_with_summaries(row: pd.Series, source: str) -> str:
    topic_count = int(pd.to_numeric(row.get(f"{source}_topic_count"), errors="coerce") or 0)
    doc_count = int(pd.to_numeric(row.get(f"{source}_document_count"), errors="coerce") or 0)
    labels = split_pipe(row.get(f"{source}_labels"))
    summaries = split_pipe(row.get(f"{source}_summaries"))
    if topic_count <= 0 or not labels:
        return "No retained counterpart."
    lines: list[str] = [f"{topic_count} topics; {doc_count} unique documents"]
    for idx, label in enumerate(labels):
        summary = paragraph(summaries[idx], 230) if idx < len(summaries) else ""
        lines.append(f"- {label}" + (f": {summary}" if summary else ""))
    return "\n".join(lines)


def source_note(values: pd.Series) -> str:
    seen: list[str] = []
    for value in values:
        for part in split_pipe(str(value).replace("; ", " | ")):
            if part and part not in seen:
                seen.append(part)
    return " | ".join(seen)


def load_tables() -> dict[str, pd.DataFrame]:
    if not FULL_XLSX.exists():
        raise FileNotFoundError(f"Run build_paper_appendix_method_tables.py first: {FULL_XLSX}")
    unpaired_frames: list[pd.DataFrame] = []
    for path in sorted(
        (
            PIPELINE_ROOT
            / "outputs"
            / "paper_tables"
            / "corporate_focus_relative_longitudinal_series"
            / "topic_tables"
        ).glob("T*_unpaired_external_topics.csv")
    ):
        frame = pd.read_csv(path)
        if frame.empty:
            continue
        frame["source_file"] = str(path.relative_to(PIPELINE_ROOT))
        unpaired_frames.append(frame)
    return {
        "macro": pd.read_excel(FULL_XLSX, sheet_name="Macro topics"),
        "corporate_map": pd.read_excel(FULL_XLSX, sheet_name="Corporate with aggregates"),
        "external": pd.read_excel(FULL_XLSX, sheet_name="External aggregate topics"),
        "unpaired": pd.concat(unpaired_frames, ignore_index=True) if unpaired_frames else pd.DataFrame(),
    }


def build_macro_overview(macro: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for row in macro.sort_values("macro_topic").itertuples(index=False):
        rows.append(
            {
                "Macro topic": f"{row.macro_topic} - {row.macro_topic_name}",
                "Analytical scope": paragraph(row.definition, 360),
                "SDG crosswalk": clean_text(row.sdg_crosswalk),
                "Assignment rule": clean_text(row.cosine_assignment_rule),
                "Boundary and exclusion notes": f"{paragraph(row.boundary_notes, 300)}\nFalse positives: {paragraph(row.false_positive_notes, 300)}",
                "Source file": clean_text(row.source_file),
            }
        )
    return pd.DataFrame(rows)


def build_expression_overview(macro: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for row in macro.sort_values("macro_topic").itertuples(index=False):
        expressions = clean_text(getattr(row, "descriptor_expressions", "")).replace("- ", "\n- ")
        rows.append(
            {
                "Macro topic": f"{row.macro_topic} - {row.macro_topic_name}",
                "Descriptor expressions embedded separately": expressions.strip(),
                "Number of expressions": int(pd.to_numeric(row.element_count, errors="coerce")),
                "How expressions were used": "Each expression was embedded as an individual subanchor; the macro-topic score is the maximum expression-level cosine similarity.",
                "Source file": clean_text(row.source_file),
            }
        )
    return pd.DataFrame(rows)


def build_compact_corporate_map(corporate_map: pd.DataFrame) -> pd.DataFrame:
    rows = []
    sort_cols = ["macro_topic", "corporate_topic_label"]
    for row in corporate_map.sort_values(sort_cols).itertuples(index=False):
        series = pd.Series(row._asdict())
        rows.append(
            {
                "Macro topic": f"{row.macro_topic} - {row.macro_topic_name}",
                "Corporate anchor": f"{row.corporate_topic_label}\n(n={int(row.corporate_document_count)} documents)",
                "Corporate summary": paragraph(row.corporate_summary, 420),
                "Academic counterparts": counterpart_cell(series, "academic"),
                "Media counterparts": counterpart_cell(series, "media"),
                "Source files": clean_text(row.source_file),
            }
        )
    return pd.DataFrame(rows)


def build_external_detail(external: pd.DataFrame) -> pd.DataFrame:
    rows = []
    sort_cols = ["macro_topic", "corporate_topic_label", "external_source", "external_topic_label"]
    for row in external.sort_values(sort_cols).itertuples(index=False):
        rows.append(
            {
                "Macro topic": f"{row.macro_topic} - {row.macro_topic_name}",
                "Corporate anchor": f"{row.corporate_topic_label}\n(n={int(row.corporate_document_count)} documents)",
                "External source": row.external_source,
                "External topic": f"{row.external_topic_label}\n(n={int(row.external_document_count)} documents)",
                "External summary": paragraph(row.external_summary, 420),
                "Review / similarity evidence": (
                    f"{clean_text(row.inclusion_reason)}; "
                    f"best cosine={pd.to_numeric(row.best_cosine_similarity, errors='coerce'):.3f}"
                ),
                "Source file": clean_text(row.source_file),
            }
        )
    return pd.DataFrame(rows)


def build_unpaired_external(unpaired: pd.DataFrame) -> pd.DataFrame:
    if unpaired.empty:
        return pd.DataFrame(
            columns=[
                "Macro topic",
                "External source",
                "Relevant unpaired external topic",
                "External summary",
                "Nearest reviewed corporate anchor",
                "Basis for review",
                "Source file",
            ]
        )
    rows = []
    sort_cols = ["macro_topic", "source_noncorporate", "topic_label_refined_noncorporate"]
    for row in unpaired.sort_values(sort_cols).itertuples(index=False):
        basis = clean_text(row.inclusion_reason)
        if basis == "direct_pair_with_corporate_above_threshold":
            basis = "Direct corporate pair above similarity threshold, but retained as relevant unpaired after substantive review"
        elif basis == "manual_override_from_excluded_review":
            basis = "Manually retained after exclusion review as relevant unpaired external signal"
        rows.append(
            {
                "Macro topic": f"{row.macro_topic} - {row.macro_topic_name}",
                "External source": row.source_noncorporate,
                "Relevant unpaired external topic": (
                    f"{row.topic_label_refined_noncorporate}\n"
                    f"(n={int(row.external_unique_document_count)} documents)"
                ),
                "External summary": paragraph(row.overall_summary_noncorporate, 460),
                "Nearest reviewed corporate anchor": (
                    f"{row.topic_label_refined_corporate}\n"
                    f"(not retained as an aligned counterpart)"
                ),
                "Basis for review": basis,
                "Source file": clean_text(row.source_file),
            }
        )
    return pd.DataFrame(rows)


def build_macro_sheet_frame(corporate_map: pd.DataFrame, macro_topic: str) -> pd.DataFrame:
    subset = corporate_map.loc[corporate_map["macro_topic"] == macro_topic].copy()
    rows = []
    for row in subset.sort_values("corporate_topic_label").itertuples(index=False):
        series = pd.Series(row._asdict())
        rows.append(
            {
                "Corporate anchor": f"{row.corporate_topic_label}\n(n={int(row.corporate_document_count)} documents)",
                "Corporate summary": paragraph(row.corporate_summary, 420),
                "Academic counterparts": counterpart_cell_with_summaries(series, "academic"),
                "Media counterparts": counterpart_cell_with_summaries(series, "media"),
            }
        )
    return pd.DataFrame(rows)


def add_title(ws, title: str, width: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(row=1, column=1, value=title)
    cell.fill = TITLE_FILL
    cell.font = WHITE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28


def write_dataframe(ws, frame: pd.DataFrame, start_row: int, start_col: int = 1) -> int:
    for col_idx, column in enumerate(frame.columns, start=start_col):
        cell = ws.cell(row=start_row, column=col_idx, value=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    for row_idx, row in enumerate(frame.itertuples(index=False), start=start_row + 1):
        fill = LIGHT_FILL if (row_idx - start_row) % 2 == 0 else None
        for col_idx, value in enumerate(row, start=start_col):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fill:
                cell.fill = fill
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
    return start_row + len(frame) + 1


def style_sheet(ws, widths: dict[int, int]) -> None:
    ws.freeze_panes = "A3"
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row in ws.iter_rows():
        max_lines = 1
        for cell in row:
            if cell.value is not None:
                max_lines = max(max_lines, str(cell.value).count("\n") + 1)
        ws.row_dimensions[row[0].row].height = min(max(18, max_lines * 14), 210)


def write_readme(wb: Workbook) -> None:
    ws = wb.create_sheet("README")
    add_title(ws, "Reader-Oriented Appendix Tables", 3)
    rows = pd.DataFrame(
        [
            {
                "Item": "Purpose",
                "Description": "This workbook is a visual companion to the full audit workbook. It is designed for appendix reading rather than data inspection.",
                "Where to use": "Use in supplementary material or to copy compact tables into the paper appendix.",
            },
            {
                "Item": "A1 Taxonomy overview",
                "Description": "Shows the six macro-topics, analytical scope, SDG crosswalk, and assignment rule.",
                "Where to use": "Method appendix for taxonomy and document-to-domain mapping.",
            },
            {
                "Item": "A2 Descriptor expressions",
                "Description": "Shows the descriptor expressions used as separate embedding subanchors for each macro-topic.",
                "Where to use": "Method appendix for the document-to-domain mapping procedure.",
            },
            {
                "Item": "A3 Corporate-external topic map",
                "Description": "Shows each retained corporate anchor and its aggregated academic/media counterparts.",
                "Where to use": "Method/results appendix for final comparative sample construction.",
            },
            {
                "Item": "A5 Relevant but unpaired external topics",
                "Description": "Shows external topics retained as substantively relevant but not treated as aligned counterparts to corporate anchors.",
                "Where to use": "Results appendix for external signals outside the corporate frame.",
            },
            {
                "Item": "T1-T6 sheets",
                "Description": "Reader-friendly domain sheets, one per macro-topic, with corporate anchors and external counterpart summaries.",
                "Where to use": "Supplemental domain tables.",
            },
            {
                "Item": "Source tracking",
                "Description": "Source files are retained in compact columns or sheet footnotes so the tables remain auditable.",
                "Where to use": "Documentation and reproducibility.",
            },
        ]
    )
    write_dataframe(ws, rows, 3)
    style_sheet(ws, {1: 24, 2: 78, 3: 50})


def write_table_sheet(wb: Workbook, sheet_name: str, title: str, frame: pd.DataFrame, widths: dict[int, int]) -> None:
    ws = wb.create_sheet(sheet_name)
    add_title(ws, title, len(frame.columns))
    write_dataframe(ws, frame, 3)
    style_sheet(ws, widths)


def write_macro_domain_sheets(wb: Workbook, macro: pd.DataFrame, corporate_map: pd.DataFrame) -> None:
    macro_lookup = macro.set_index("macro_topic")
    for macro_topic in MACRO_ORDER:
        if macro_topic not in macro_lookup.index:
            continue
        row = macro_lookup.loc[macro_topic]
        title = f"{macro_topic} - {row['macro_topic_name']}"
        frame = build_macro_sheet_frame(corporate_map, macro_topic)
        ws = wb.create_sheet(macro_topic)
        add_title(ws, title, 4)
        ws.cell(row=3, column=1, value="Analytical scope")
        ws.cell(row=3, column=2, value=paragraph(row["definition"], 600))
        ws.cell(row=4, column=1, value="Embedding rule")
        ws.cell(row=4, column=2, value=clean_text(row["cosine_assignment_rule"]))
        ws.cell(row=5, column=1, value="Descriptor expressions")
        ws.cell(row=5, column=2, value=clean_text(row["descriptor_expressions"]).replace("- ", "\n- ").strip())
        for r in [3, 4, 5]:
            ws.cell(row=r, column=1).fill = SECTION_FILL
            ws.cell(row=r, column=1).font = HEADER_FONT
            ws.cell(row=r, column=2).font = BODY_FONT
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        write_dataframe(ws, frame, 7)
        style_sheet(ws, {1: 34, 2: 58, 3: 68, 4: 68})


def markdown_table(frame: pd.DataFrame, columns: list[str]) -> str:
    trimmed = frame[columns].copy()
    for column in trimmed.columns:
        trimmed[column] = trimmed[column].map(
            lambda value: "<br>".join(
                " ".join(line.split()).replace("|", "\\|")
                for line in str(value if value is not None and not pd.isna(value) else "").splitlines()
                if " ".join(line.split())
            )
        )
    header = "| " + " | ".join(trimmed.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(trimmed.columns)) + " |"
    rows = [
        "| " + " | ".join(str(value) for value in row) + " |"
        for row in trimmed.itertuples(index=False, name=None)
    ]
    return "\n".join([header, separator, *rows])


def write_markdown(
    macro_overview: pd.DataFrame,
    expression_overview: pd.DataFrame,
    compact_map: pd.DataFrame,
    unpaired_external: pd.DataFrame,
) -> None:
    parts = [
        "# Paper Appendix Visual Tables",
        "",
        "These tables are reader-oriented renderings of the full appendix workbook. Source files are retained in the Excel version.",
        "",
        "## Table A1. Macro-topic taxonomy",
        "",
        markdown_table(
            macro_overview,
            [
                "Macro topic",
                "Analytical scope",
                "SDG crosswalk",
                "Assignment rule",
                "Boundary and exclusion notes",
            ],
        ),
        "",
        "## Table A2. Descriptor expressions by macro-topic",
        "",
        markdown_table(
            expression_overview,
            [
                "Macro topic",
                "Descriptor expressions embedded separately",
                "Number of expressions",
                "How expressions were used",
            ],
        ),
        "",
        "## Table A3. Corporate anchors and external counterparts",
        "",
        markdown_table(
            compact_map,
            [
                "Macro topic",
                "Corporate anchor",
                "Corporate summary",
                "Academic counterparts",
                "Media counterparts",
            ],
        ),
        "",
        "## Table A5. Relevant external topics without corporate counterparts",
        "",
        markdown_table(
            unpaired_external,
            [
                "Macro topic",
                "External source",
                "Relevant unpaired external topic",
                "External summary",
                "Nearest reviewed corporate anchor",
                "Basis for review",
            ],
        ),
        "",
    ]
    OUTPUT_MD.write_text("\n".join(parts), encoding="utf-8")


def main() -> None:
    APPENDIX_DIR.mkdir(parents=True, exist_ok=True)
    tables = load_tables()
    macro_overview = build_macro_overview(tables["macro"])
    expression_overview = build_expression_overview(tables["macro"])
    compact_map = build_compact_corporate_map(tables["corporate_map"])
    external_detail = build_external_detail(tables["external"])
    unpaired_external = build_unpaired_external(tables["unpaired"])

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    write_readme(wb)
    write_table_sheet(
        wb,
        "A1 Macro Overview",
        "Appendix Table A1. Environmental Macro-Topic Taxonomy",
        macro_overview,
        {1: 32, 2: 52, 3: 18, 4: 56, 5: 62, 6: 44},
    )
    write_table_sheet(
        wb,
        "A2 Expressions",
        "Appendix Table A2. Descriptor Expressions by Macro-Topic",
        expression_overview,
        {1: 32, 2: 70, 3: 18, 4: 58, 5: 44},
    )
    write_table_sheet(
        wb,
        "A3 Corporate Map",
        "Appendix Table A3. Corporate Anchors and External Counterparts",
        compact_map,
        {1: 32, 2: 42, 3: 58, 4: 62, 5: 62, 6: 48},
    )
    write_table_sheet(
        wb,
        "A4 External Details",
        "Appendix Table A4. External Topics Retained as Aligned Counterparts",
        external_detail,
        {1: 30, 2: 42, 3: 18, 4: 46, 5: 64, 6: 34, 7: 42},
    )
    write_table_sheet(
        wb,
        "A5 Unpaired External",
        "Appendix Table A5. Relevant External Topics Without Corporate Counterparts",
        unpaired_external,
        {1: 32, 2: 18, 3: 46, 4: 64, 5: 44, 6: 48, 7: 42},
    )
    write_macro_domain_sheets(wb, tables["macro"], tables["corporate_map"])
    wb.save(OUTPUT_XLSX)
    write_markdown(macro_overview, expression_overview, compact_map, unpaired_external)
    print(f"Wrote {OUTPUT_XLSX}")
    print(f"Wrote {OUTPUT_MD}")


if __name__ == "__main__":
    main()
