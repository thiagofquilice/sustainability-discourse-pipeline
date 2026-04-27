#!/usr/bin/env python3
"""Build an Excel workbook comparing temporal narratives for correlated source-topic pairs."""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path
from typing import Any

import pandas as pd


PIPELINE_ROOT = Path("paper_pipeline")
REVIEW_ROOT = PIPELINE_ROOT / "outputs" / "corporate_focus_review_with_overrides"
RELATION_ROOT = PIPELINE_ROOT / "outputs" / "paper_tables" / "corporate_focus_source_topic_relations"
OUTPUT_XLSX = RELATION_ROOT / "source_topic_temporal_evolution_comparisons.xlsx"

GROUPS_PATH = REVIEW_ROOT / "corporate_focus_all_groups_optional.csv"
AGGREGATE_SUMMARY_PATH = RELATION_ROOT / "aggregate_source_relations_summary.csv"
INDIVIDUAL_SUMMARY_PATH = RELATION_ROOT / "individual_source_relations_summary.csv"


PHASE_COLUMNS = [
    "phase_1_years",
    "phase_1_summary",
    "phase_2_years",
    "phase_2_summary",
    "phase_3_years",
    "phase_3_summary",
    "phase_4_years",
    "phase_4_summary",
]


def clean_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return " ".join(str(value).split())


def short_text(value: object, max_chars: int = 420) -> str:
    text = clean_text(value)
    if len(text) <= max_chars:
        return text
    return text[:max_chars].rsplit(" ", 1)[0].strip() + "..."


def safe_float(value: object) -> float:
    converted = pd.to_numeric(value, errors="coerce")
    return float(converted) if pd.notna(converted) else float("nan")


def safe_int_text(value: object) -> str:
    converted = pd.to_numeric(value, errors="coerce")
    if pd.isna(converted):
        return "not available"
    return str(int(converted))


def lag_phrase(lag: object) -> str:
    converted = pd.to_numeric(lag, errors="coerce")
    if pd.isna(converted):
        return "no valid lag"
    lag_int = int(converted)
    if lag_int > 0:
        return f"external discourse leading corporate disclosure by {lag_int} year(s)"
    if lag_int < 0:
        return f"corporate disclosure leading external discourse by {abs(lag_int)} year(s)"
    return "same-year coevolution"


def relation_phrase(label: object, lag: object, r_value: object) -> str:
    label_text = clean_text(label)
    r = safe_float(r_value)
    sign = "positive" if pd.notna(r) and r >= 0 else "negative"
    if label_text == "external_leads":
        return f"an apparent external lead, with a {sign} association at the selected lag"
    if label_text == "corporate_leads":
        return f"an apparent corporate lead, with a {sign} association at the selected lag"
    return f"a synchronous or ambiguous relation, with a {sign} association at the selected lag"


def load_topic_metadata() -> pd.DataFrame:
    groups = pd.read_csv(GROUPS_PATH)
    groups["final_merge_group_id"] = groups["final_merge_group_id"].astype(str)
    groups["topic_key"] = groups["subgroup"].astype(str) + "::" + groups["final_merge_group_id"].astype(str)
    for column in ["topic_label_refined", "topic_name_original", "overall_summary", "evolution_pattern", *PHASE_COLUMNS]:
        if column not in groups.columns:
            groups[column] = ""
    return groups


def metadata_lookup(groups: pd.DataFrame) -> dict[str, dict[str, Any]]:
    return {
        str(row.topic_key): row._asdict()
        for row in groups.itertuples(index=False)
    }


def topic_key(subgroup: object, final_merge_group_id: object) -> str:
    return f"{subgroup}::{final_merge_group_id}"


def topic_meta(lookup: dict[str, dict[str, Any]], subgroup: object, final_merge_group_id: object) -> dict[str, Any]:
    return lookup.get(topic_key(subgroup, str(final_merge_group_id)), {})


def topic_label(meta: dict[str, Any], fallback: object = "") -> str:
    return clean_text(meta.get("topic_label_refined", "")) or clean_text(meta.get("topic_name_original", "")) or clean_text(fallback)


def phases_compact(meta: dict[str, Any], max_phase_chars: int = 220) -> str:
    phases: list[str] = []
    for idx in range(1, 5):
        years = clean_text(meta.get(f"phase_{idx}_years", ""))
        summary = short_text(meta.get(f"phase_{idx}_summary", ""), max_phase_chars)
        if years or summary:
            phases.append(f"Phase {idx} ({years or 'years not specified'}): {summary}")
    return "\n".join(phases)


def add_topic_description_columns(row: dict[str, Any], prefix: str, meta: dict[str, Any]) -> None:
    row[f"{prefix}_topic_label"] = topic_label(meta, row.get(f"{prefix}_topic_label", ""))
    row[f"{prefix}_overall_summary"] = clean_text(meta.get("overall_summary", ""))
    row[f"{prefix}_evolution_pattern"] = clean_text(meta.get("evolution_pattern", ""))
    for idx in range(1, 5):
        row[f"{prefix}_phase_{idx}_years"] = clean_text(meta.get(f"phase_{idx}_years", ""))
        row[f"{prefix}_phase_{idx}_summary"] = clean_text(meta.get(f"phase_{idx}_summary", ""))
    row[f"{prefix}_phases_compact"] = phases_compact(meta)


def aggregate_external_details(
    individual: pd.DataFrame,
    lookup: dict[str, dict[str, Any]],
    macro_topic: object,
    corporate_group_id: object,
    external_source: object,
) -> tuple[str, str, str]:
    subset = individual[
        (individual["macro_topic"].astype(str) == str(macro_topic))
        & (individual["corporate_final_merge_group_id"].astype(str) == str(corporate_group_id))
        & (individual["external_source"].astype(str) == str(external_source))
    ].copy()
    if subset.empty:
        return "", "", ""
    subset = subset.sort_values("external_unique_document_count", ascending=False, kind="mergesort")
    labels: list[str] = []
    summaries: list[str] = []
    phases: list[str] = []
    for relation in subset.itertuples(index=False):
        meta = topic_meta(lookup, relation.external_subgroup, relation.external_final_merge_group_id)
        label = topic_label(meta, getattr(relation, "external_topic_label", ""))
        labels.append(label)
        summary = short_text(meta.get("overall_summary", ""), 260)
        if summary:
            summaries.append(f"{label}: {summary}")
        phase_text = phases_compact(meta, max_phase_chars=140)
        if phase_text:
            phases.append(f"{label}\n{phase_text}")
    return "; ".join(labels), "\n\n".join(summaries), "\n\n".join(phases)


def build_comparison_paragraph(row: pd.Series) -> str:
    corporate_label = clean_text(row.get("corporate_topic_label", "the corporate topic"))
    external_label = clean_text(row.get("external_topic_label", "")) or clean_text(
        row.get("external_relation_label", "the external counterpart")
    )
    source = clean_text(row.get("external_source", "external"))
    corporate_summary = short_text(row.get("corporate_overall_summary", ""), 360)
    external_summary = short_text(row.get("external_overall_summary", ""), 360)
    corporate_phase = short_text(row.get("corporate_phases_compact", ""), 360)
    external_phase = short_text(row.get("external_phases_compact", ""), 360)
    lag = row.get("best_lag_spearman", "")
    r_value = row.get("best_spearman_r", "")
    q_value = row.get("best_spearman_q_table", "")
    relation = relation_phrase(row.get("temporal_relation_label", ""), lag, r_value)
    timing = lag_phrase(lag)
    first_gap = safe_int_text(row.get("first_active_year_gap", ""))
    peak_gap = safe_int_text(row.get("peak_year_gap", ""))
    q_text = "not available" if pd.isna(pd.to_numeric(q_value, errors="coerce")) else f"{float(q_value):.3g}"
    r_text = "not available" if pd.isna(pd.to_numeric(r_value, errors="coerce")) else f"{float(r_value):.3f}"

    return (
        f"The corporate topic \"{corporate_label}\" and the {source} counterpart \"{external_label}\" show {relation}; "
        f"the strongest Spearman association is at lag {safe_int_text(lag)} ({timing}), with r={r_text} and table-level q={q_text}. "
        f"The corporate narrative is summarized as: {corporate_summary or 'no corporate narrative summary available'}. "
        f"Its period narrative highlights {corporate_phase or 'no corporate phase description available'}. "
        f"The {source} narrative is summarized as: {external_summary or 'no external narrative summary available'}. "
        f"Its period narrative highlights {external_phase or 'no external phase description available'}. "
        f"The first-active-year gap is {first_gap} year(s) and the peak-year gap is {peak_gap} year(s), so the relation is best read as an exploratory temporal alignment rather than causal evidence."
    )


def select_base_columns(frame: pd.DataFrame) -> list[str]:
    preferred = [
        "relation_id",
        "relation_level",
        "macro_topic",
        "macro_topic_name",
        "corporate_subgroup",
        "corporate_final_merge_group_id",
        "corporate_topic_label",
        "external_source",
        "external_subgroup",
        "external_final_merge_group_id",
        "external_topic_label",
        "external_relation_label",
        "external_topic_count",
        "corporate_unique_document_count",
        "external_unique_document_count",
        "best_lag_spearman",
        "best_spearman_r",
        "best_spearman_p",
        "best_spearman_q_relation",
        "best_spearman_q_table",
        "pearson_r_at_best_spearman_lag",
        "best_overlap_years",
        "first_active_year_gap",
        "peak_year_gap",
        "temporal_relation_label",
        "temporal_relation_reason",
        "ambiguous_best_lag_flag",
        "corporate_first_active_year",
        "external_first_active_year",
        "corporate_peak_year",
        "external_peak_year",
    ]
    return [column for column in preferred if column in frame.columns]


def build_individual_comparisons(individual: pd.DataFrame, lookup: dict[str, dict[str, Any]]) -> pd.DataFrame:
    valid = individual[pd.to_numeric(individual["best_spearman_r"], errors="coerce").notna()].copy()
    rows: list[dict[str, Any]] = []
    for relation in valid.itertuples(index=False):
        base = {column: getattr(relation, column) for column in valid.columns}
        corporate_meta = topic_meta(lookup, relation.corporate_subgroup, relation.corporate_final_merge_group_id)
        external_meta = topic_meta(lookup, relation.external_subgroup, relation.external_final_merge_group_id)
        add_topic_description_columns(base, "corporate", corporate_meta)
        add_topic_description_columns(base, "external", external_meta)
        rows.append(base)
    result = pd.DataFrame(rows)
    result["comparison_paragraph"] = result.apply(build_comparison_paragraph, axis=1)
    ordered = select_base_columns(result) + [
        "corporate_overall_summary",
        "corporate_evolution_pattern",
        "corporate_phase_1_years",
        "corporate_phase_1_summary",
        "corporate_phase_2_years",
        "corporate_phase_2_summary",
        "corporate_phase_3_years",
        "corporate_phase_3_summary",
        "corporate_phase_4_years",
        "corporate_phase_4_summary",
        "external_overall_summary",
        "external_evolution_pattern",
        "external_phase_1_years",
        "external_phase_1_summary",
        "external_phase_2_years",
        "external_phase_2_summary",
        "external_phase_3_years",
        "external_phase_3_summary",
        "external_phase_4_years",
        "external_phase_4_summary",
        "comparison_paragraph",
    ]
    ordered += [column for column in result.columns if column not in ordered]
    return result[ordered]


def build_aggregate_comparisons(
    aggregate: pd.DataFrame,
    individual: pd.DataFrame,
    lookup: dict[str, dict[str, Any]],
) -> pd.DataFrame:
    valid = aggregate[pd.to_numeric(aggregate["best_spearman_r"], errors="coerce").notna()].copy()
    rows: list[dict[str, Any]] = []
    for relation in valid.itertuples(index=False):
        base = {column: getattr(relation, column) for column in valid.columns}
        corporate_meta = topic_meta(lookup, relation.corporate_subgroup, relation.corporate_final_merge_group_id)
        add_topic_description_columns(base, "corporate", corporate_meta)
        labels, summaries, phases = aggregate_external_details(
            individual,
            lookup,
            relation.macro_topic,
            relation.corporate_final_merge_group_id,
            relation.external_source,
        )
        base["external_topic_label"] = f"{relation.external_source} aggregate"
        base["external_component_labels"] = labels
        base["external_overall_summary"] = summaries
        base["external_evolution_pattern"] = ""
        base["external_phases_compact"] = phases
        for idx in range(1, 5):
            base[f"external_phase_{idx}_years"] = ""
            base[f"external_phase_{idx}_summary"] = ""
        rows.append(base)
    result = pd.DataFrame(rows)
    result["comparison_paragraph"] = result.apply(build_comparison_paragraph, axis=1)
    ordered = select_base_columns(result) + [
        "corporate_overall_summary",
        "corporate_evolution_pattern",
        "corporate_phase_1_years",
        "corporate_phase_1_summary",
        "corporate_phase_2_years",
        "corporate_phase_2_summary",
        "corporate_phase_3_years",
        "corporate_phase_3_summary",
        "corporate_phase_4_years",
        "corporate_phase_4_summary",
        "external_component_labels",
        "external_overall_summary",
        "external_phases_compact",
        "comparison_paragraph",
    ]
    ordered += [column for column in result.columns if column not in ordered]
    return result[ordered]


def autosize_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:80])
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 55)
        for row in worksheet.iter_rows():
            for cell in row:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "top"
                cell.alignment = alignment
    workbook.save(path)


def main() -> None:
    groups = load_topic_metadata()
    lookup = metadata_lookup(groups)
    aggregate = pd.read_csv(AGGREGATE_SUMMARY_PATH)
    individual = pd.read_csv(INDIVIDUAL_SUMMARY_PATH)

    aggregate_comparisons = build_aggregate_comparisons(aggregate, individual, lookup)
    individual_comparisons = build_individual_comparisons(individual, lookup)

    readme = pd.DataFrame(
        [
            {
                "item": "unit",
                "description": "Rows are corporate-external topic relations with a valid best Spearman correlation.",
            },
            {
                "item": "metric",
                "description": "Temporal association uses annual_document_prevalence, i.e. unique topic documents in a year divided by source-domain documents in that year.",
            },
            {
                "item": "lag convention",
                "description": "Positive lags mean academic/media external discourse leads corporate disclosure; negative lags mean corporate leads external discourse.",
            },
            {
                "item": "interpretation",
                "description": "The generated paragraph is an interpretive synthesis of relation diagnostics plus the stored temporal narrative fields. It is exploratory and not causal.",
            },
            {
                "item": "row counts",
                "description": json.dumps(
                    {
                        "aggregate_comparisons": int(len(aggregate_comparisons)),
                        "individual_comparisons": int(len(individual_comparisons)),
                    },
                    ensure_ascii=False,
                ),
            },
        ]
    )

    RELATION_ROOT.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        aggregate_comparisons.to_excel(writer, sheet_name="Aggregate comparisons", index=False)
        individual_comparisons.to_excel(writer, sheet_name="Individual comparisons", index=False)

    autosize_workbook(OUTPUT_XLSX)
    print(
        json.dumps(
            {
                "output": str(OUTPUT_XLSX),
                "aggregate_comparisons": int(len(aggregate_comparisons)),
                "individual_comparisons": int(len(individual_comparisons)),
            },
            indent=2,
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
