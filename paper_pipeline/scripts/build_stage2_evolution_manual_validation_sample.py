#!/usr/bin/env python3
"""Build a manual validation sample for Stage 2 micro-topic evolution narratives."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd


WORKFLOW_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = WORKFLOW_ROOT / "outputs" / "bertopic_micro_evolution"
DEFAULT_NARRATIVES = OUTPUT_ROOT / "evolution_summaries" / "micro_topic_evolution_narratives.csv"
DEFAULT_ANNUAL = OUTPUT_ROOT / "year_summaries" / "micro_topic_year_summaries.csv"
DEFAULT_SELECTED = OUTPUT_ROOT / "selected_micro_topics.csv"
DEFAULT_OUTPUT_DIR = OUTPUT_ROOT / "manual_validation_sample"

CORE_FIELDS = [
    "topic_label_refined",
    "overall_summary",
    "phase_1_years",
    "phase_1_summary",
    "phase_2_years",
    "phase_2_summary",
    "phase_3_years",
    "phase_3_summary",
    "phase_4_years",
    "phase_4_summary",
    "evolution_pattern",
    "evidence_note",
]

MANUAL_COLUMNS = [
    "manual_overall_judgment",
    "manual_label_quality",
    "manual_phase_structure_quality",
    "manual_evidence_support",
    "manual_blank_fields_problem",
    "manual_notes",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--narratives", type=Path, default=DEFAULT_NARRATIVES)
    parser.add_argument("--annual-summaries", type=Path, default=DEFAULT_ANNUAL)
    parser.add_argument("--selected-topics", type=Path, default=DEFAULT_SELECTED)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--seed", type=int, default=42)
    return parser.parse_args()


def normalize_text(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if text.lower() == "none":
        return ""
    return text


def build_annual_digest(group: pd.DataFrame) -> dict[str, str | int]:
    group = group.sort_values("year").copy()
    lines = []
    for _, row in group.iterrows():
        year = int(row["year"])
        primary = normalize_text(row.get("year_focus_primary", ""))
        secondary = normalize_text(row.get("year_focus_secondary", ""))
        frame = normalize_text(row.get("year_frame_or_angle", ""))
        evidence = normalize_text(row.get("year_evidence_note", ""))
        line = f"{year} | primary={primary}"
        if secondary:
            line += f" | secondary={secondary}"
        if frame:
            line += f" | frame={frame}"
        if evidence:
            line += f" | note={evidence}"
        lines.append(line)
    years = group["year"].astype(int).tolist()
    return {
        "annual_summary_count": int(len(group)),
        "annual_first_year": int(min(years)) if years else pd.NA,
        "annual_last_year": int(max(years)) if years else pd.NA,
        "annual_years_covered": ", ".join(str(year) for year in years),
        "annual_summary_digest": "\n".join(lines),
    }


def add_completeness_fields(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in CORE_FIELDS:
        result[column] = result[column].map(normalize_text)

    present_counts = []
    missing_counts = []
    missing_lists = []
    phase_counts = []
    has_core_blank = []
    for _, row in result.iterrows():
        missing = [column for column in CORE_FIELDS if not normalize_text(row[column])]
        present = len(CORE_FIELDS) - len(missing)
        phase_present = sum(
            1
            for column in ["phase_1_years", "phase_2_years", "phase_3_years", "phase_4_years"]
            if normalize_text(row[column])
        )
        present_counts.append(present)
        missing_counts.append(len(missing))
        missing_lists.append(" | ".join(missing))
        phase_counts.append(phase_present)
        has_core_blank.append(bool(missing))

    result["core_fields_present_count"] = present_counts
    result["core_fields_missing_count"] = missing_counts
    result["missing_core_fields"] = missing_lists
    result["phase_count_present"] = phase_counts
    result["has_core_blank"] = has_core_blank
    return result


def select_sample(frame: pd.DataFrame, seed: int) -> pd.DataFrame:
    selected_parts: list[pd.DataFrame] = []
    ordered = frame.sort_values(["subgroup", "micro_topic_id"]).copy()
    for subgroup, subgroup_frame in ordered.groupby("subgroup", sort=True):
        subgroup_frame = subgroup_frame.sample(frac=1.0, random_state=seed).copy()

        most_complete = subgroup_frame.sort_values(
            ["core_fields_present_count", "phase_count_present", "annual_summary_count", "micro_topic_id"],
            ascending=[False, False, False, True],
            kind="mergesort",
        ).head(1).copy()
        most_complete["sample_reason"] = "most_complete_in_subgroup"
        selected_parts.append(most_complete)

        remaining = subgroup_frame.loc[
            ~subgroup_frame["micro_topic_id"].isin(most_complete["micro_topic_id"].tolist())
        ].copy()
        if not remaining.empty:
            least_complete = remaining.sort_values(
                ["core_fields_present_count", "phase_count_present", "annual_summary_count", "micro_topic_id"],
                ascending=[True, True, True, True],
                kind="mergesort",
            ).head(1).copy()
            least_complete["sample_reason"] = "least_complete_in_subgroup"
            selected_parts.append(least_complete)

    sample = pd.concat(selected_parts, ignore_index=True)
    sample = sample.sort_values(["subgroup", "sample_reason", "micro_topic_id"]).reset_index(drop=True)
    sample["review_id"] = [f"S2_{i + 1:03d}" for i in range(len(sample))]
    for column in MANUAL_COLUMNS:
        sample[column] = pd.NA
    return sample


def set_workbook_formatting(xlsx_path: Path) -> None:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except Exception:
        return

    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    width_overrides = {
        "A": 12,
        "B": 14,
        "C": 12,
        "D": 10,
        "E": 14,
        "F": 36,
        "G": 18,
        "H": 10,
        "I": 10,
        "J": 14,
        "K": 22,
        "L": 16,
        "M": 16,
        "N": 16,
        "O": 16,
        "P": 18,
        "Q": 28,
        "R": 28,
        "S": 18,
        "T": 40,
        "U": 16,
        "V": 16,
        "W": 16,
        "X": 50,
        "Y": 40,
        "Z": 40,
        "AA": 20,
        "AB": 20,
        "AC": 20,
        "AD": 20,
        "AE": 20,
        "AF": 40,
    }
    for idx, _ in enumerate(sheet[1], start=1):
        col = get_column_letter(idx)
        sheet.column_dimensions[col].width = width_overrides.get(col, 18)

    workbook.save(xlsx_path)


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    narratives = pd.read_csv(args.narratives)
    annual = pd.read_csv(args.annual_summaries)
    selected = pd.read_csv(args.selected_topics)

    annual_grouped = (
        annual.groupby(["subgroup", "micro_topic_id"], dropna=False)
        .apply(build_annual_digest)
        .apply(pd.Series)
        .reset_index()
    )

    keep_selected = [
        "subgroup",
        "source",
        "assigned_label",
        "micro_topic_id",
        "topic_name_original",
        "topic_size",
        "share_of_non_outlier",
        "overall_keywords",
    ]
    merged = selected[keep_selected].merge(
        narratives,
        on=["subgroup", "source", "assigned_label", "micro_topic_id", "topic_name_original"],
        how="left",
        validate="one_to_one",
    ).merge(
        annual_grouped,
        on=["subgroup", "micro_topic_id"],
        how="left",
        validate="one_to_one",
    )

    merged = add_completeness_fields(merged)
    sample = select_sample(merged, args.seed)

    ordered_columns = [
        "review_id",
        "sample_reason",
        "subgroup",
        "source",
        "assigned_label",
        "micro_topic_id",
        "topic_name_original",
        "topic_size",
        "share_of_non_outlier",
        "overall_keywords",
        "core_fields_present_count",
        "core_fields_missing_count",
        "missing_core_fields",
        "phase_count_present",
        "has_core_blank",
        "topic_label_refined",
        "overall_summary",
        "phase_1_years",
        "phase_1_summary",
        "phase_2_years",
        "phase_2_summary",
        "phase_3_years",
        "phase_3_summary",
        "phase_4_years",
        "phase_4_summary",
        "evolution_pattern",
        "evidence_note",
        "annual_summary_count",
        "annual_first_year",
        "annual_last_year",
        "annual_years_covered",
        "annual_summary_digest",
        *MANUAL_COLUMNS,
    ]
    sample = sample[ordered_columns]

    csv_path = args.output_dir / "stage2_evolution_manual_validation_sample.csv"
    xlsx_path = args.output_dir / "stage2_evolution_manual_validation_sample.xlsx"
    manifest_path = args.output_dir / "stage2_evolution_manual_validation_sample_manifest.json"
    readme_path = args.output_dir / "stage2_evolution_manual_validation_sample_README.md"

    sample.to_csv(csv_path, index=False)
    sample.to_excel(xlsx_path, index=False)
    set_workbook_formatting(xlsx_path)

    manifest = {
        "rows": int(len(sample)),
        "sampling_rule": "two_per_subgroup_most_and_least_complete_when_available",
        "seed": args.seed,
        "source_narratives": str(args.narratives),
        "source_annual_summaries": str(args.annual_summaries),
        "source_selected_topics": str(args.selected_topics),
        "subgroup_counts": {str(k): int(v) for k, v in sample["subgroup"].value_counts().sort_index().to_dict().items()},
        "sample_reason_counts": {str(k): int(v) for k, v in sample["sample_reason"].value_counts().sort_index().to_dict().items()},
        "blank_core_rows_in_full_stage2": int(merged["has_core_blank"].sum()),
        "full_stage2_rows": int(len(merged)),
    }
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    readme_lines = [
        "# Stage 2 Manual Validation Sample",
        "",
        "This sample is for manual review of Stage 2 micro-topic evolution narratives.",
        "",
        "Sampling rule:",
        "- deterministic sample with random_state=42",
        "- up to 2 rows per subgroup",
        "- one `most_complete_in_subgroup` row",
        "- one `least_complete_in_subgroup` row",
        "",
        "Suggested review columns:",
        "- `manual_overall_judgment`: valid / partly_valid / invalid",
        "- `manual_label_quality`: good / partial / poor",
        "- `manual_phase_structure_quality`: good / partial / poor",
        "- `manual_evidence_support`: strong / partial / weak",
        "- `manual_blank_fields_problem`: none / mild / strong",
        "- `manual_notes`: free text",
        "",
        "Helpful diagnostic columns:",
        "- `core_fields_present_count`",
        "- `core_fields_missing_count`",
        "- `missing_core_fields`",
        "- `phase_count_present`",
        "- `annual_summary_digest`",
    ]
    readme_path.write_text("\n".join(readme_lines), encoding="utf-8")

    print(csv_path)


if __name__ == "__main__":
    main()
