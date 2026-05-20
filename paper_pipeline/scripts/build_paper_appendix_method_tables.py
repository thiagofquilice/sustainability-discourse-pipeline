#!/usr/bin/env python3
"""Build Excel appendix tables for the paper method supplement."""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path

import pandas as pd


PIPELINE_ROOT = Path("paper_pipeline")
CATALOG_CSV = PIPELINE_ROOT / "catalog" / "six_topic_discourse_catalog.csv"
CORPORATE_GROUPS_CSV = (
    PIPELINE_ROOT
    / "outputs"
    / "corporate_focus_review_with_overrides"
    / "corporate_focus_all_groups_optional.csv"
)
LONGITUDINAL_SERIES_CSV = (
    PIPELINE_ROOT
    / "outputs"
    / "paper_tables"
    / "corporate_focus_relative_longitudinal_series"
    / "corporate_topic_external_relative_longitudinal.csv"
)
TOPIC_TABLES_DIR = (
    PIPELINE_ROOT
    / "outputs"
    / "paper_tables"
    / "corporate_focus_relative_longitudinal_series"
    / "topic_tables"
)
OUTPUT_DIR = PIPELINE_ROOT / "outputs" / "paper_tables" / "appendix_method_tables"
OUTPUT_XLSX = OUTPUT_DIR / "paper_appendix_method_tables.xlsx"

MACRO_TOPIC_ORDER = ["T1", "T2", "T3", "T4", "T5", "T6"]
MACRO_TOPIC_NAMES = {
    "T1": "Clean energy transition",
    "T2": "Operational sustainability and circular production",
    "T3": "Sustainable products, services and consumption",
    "T4": "Climate strategy, carbon governance and disclosure",
    "T5": "Climate risk, adaptation and resilience",
    "T6": "Ecosystems, pollution and environmental stewardship",
}


def source_label(path: Path) -> str:
    return str(path.relative_to(PIPELINE_ROOT))


def clean_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return " ".join(str(value).split())


def multiline_clean(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    lines = [line.strip() for line in str(value).splitlines()]
    return "\n".join(line for line in lines if line)


def split_expressions(value: object) -> list[str]:
    expressions: list[str] = []
    for line in multiline_clean(value).splitlines():
        expression = line.strip().lstrip("-").strip()
        if expression:
            expressions.append(expression)
    return expressions


def join_values(values: pd.Series) -> str:
    seen: set[str] = set()
    output: list[str] = []
    for value in values:
        text = clean_text(value)
        if text and text not in seen:
            output.append(text)
            seen.add(text)
    return " | ".join(output)


def numeric_max(values: pd.Series) -> int:
    converted = pd.to_numeric(values, errors="coerce").dropna()
    if converted.empty:
        return 0
    return int(converted.max())


def numeric_sum(values: pd.Series) -> int:
    converted = pd.to_numeric(values, errors="coerce").dropna()
    if converted.empty:
        return 0
    return int(converted.sum())


def build_readme() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "item": "workbook purpose",
                "description": "Appendix-ready method tables for the paper's corporate-centered sustainability discourse analysis.",
            },
            {
                "item": "macro topic tables",
                "description": "Macro-topic definitions and descriptor expressions used for source-domain mapping. Cosine assignment embedded each descriptor expression separately and used the maximum expression-level score within each macro-topic.",
            },
            {
                "item": "corporate/external tables",
                "description": "Corporate anchors and the academic/media topics retained as aligned counterparts after final review.",
            },
            {
                "item": "document counts",
                "description": "Document-count fields use unique-document counts already present in final paper artifacts; raw document-topic files are not read.",
            },
            {
                "item": "source tracking",
                "description": "Every data sheet includes source_file, identifying the artifact(s) from which the row was derived.",
            },
        ]
    )


def build_sources(topic_table_paths: list[Path]) -> pd.DataFrame:
    rows = [
        {
            "sheet": "Macro topics; Macro expressions long",
            "source_file": source_label(CATALOG_CSV),
            "description": "Six macro-topic catalog with definitions, descriptor expressions, descriptor summary text, boundary notes, and false-positive notes. The operational cosine embedding units are the individual expressions listed in Macro expressions long.",
        },
        {
            "sheet": "Corporate anchors",
            "source_file": source_label(CORPORATE_GROUPS_CSV),
            "description": "Reviewed corporate and non-corporate group metadata; corporate rows provide labels, summaries, and merge metadata.",
        },
        {
            "sheet": "Corporate anchors; Corporate with aggregates",
            "source_file": source_label(LONGITUDINAL_SERIES_CSV),
            "description": "Final longitudinal series; provides unique-document counts for corporate anchors and aggregate external counterparts.",
        },
    ]
    for path in topic_table_paths:
        rows.append(
            {
                "sheet": "Corporate with aggregates; External aggregate topics",
                "source_file": source_label(path),
                "description": "Final aligned academic/media topics for one macro-topic, including labels, summaries, document counts, similarity metadata, and corporate anchor links.",
            }
        )
    return pd.DataFrame(rows)


def build_macro_topics(catalog: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "topic_code",
        "topic_name",
        "topic_definition",
        "topic_elements_bullets",
        "topic_summary_text",
        "sdg_crosswalk",
        "boundary_notes",
        "false_positive_notes",
        "element_count",
    ]
    result = catalog[columns].copy()
    result = result.rename(
        columns={
            "topic_code": "macro_topic",
            "topic_name": "macro_topic_name",
            "topic_definition": "definition",
            "topic_elements_bullets": "descriptor_expressions",
            "topic_summary_text": "descriptor_summary_text",
        }
    )
    result["descriptor_expressions"] = result["descriptor_expressions"].map(multiline_clean)
    result["cosine_embedding_unit"] = "individual descriptor expression"
    result[
        "cosine_assignment_rule"
    ] = "Each descriptor expression was embedded separately; the macro-topic score is the maximum cosine similarity across its expressions."
    result["source_file"] = source_label(CATALOG_CSV)
    return result


def build_macro_expressions(catalog: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for row in catalog.itertuples(index=False):
        for order, expression in enumerate(split_expressions(row.topic_elements_bullets), start=1):
            rows.append(
                {
                    "macro_topic": row.topic_code,
                    "macro_topic_name": row.topic_name,
                    "expression_order": order,
                    "expression": expression,
                    "cosine_embedding_unit": "yes",
                    "assignment_rule": "Embedded as one subanchor; macro-topic score uses the maximum subanchor similarity within this macro-topic.",
                    "source_file": source_label(CATALOG_CSV),
                }
            )
    return pd.DataFrame(rows)


def load_aligned_topics() -> tuple[pd.DataFrame, list[Path]]:
    frames: list[pd.DataFrame] = []
    paths: list[Path] = []
    for macro_topic in MACRO_TOPIC_ORDER:
        path = TOPIC_TABLES_DIR / f"{macro_topic}_aligned_aggregate_topics.csv"
        if not path.exists():
            continue
        frame = pd.read_csv(path)
        if frame.empty:
            continue
        frame["source_file"] = source_label(path)
        frames.append(frame)
        paths.append(path)
    if not frames:
        return pd.DataFrame(), paths
    aligned = pd.concat(frames, ignore_index=True)
    for column in [
        "final_merge_group_id_corporate",
        "final_merge_group_id_noncorporate",
        "subgroup_corporate",
        "subgroup_noncorporate",
    ]:
        if column in aligned.columns:
            aligned[column] = aligned[column].astype(str)
    return aligned, paths


def load_corporate_document_counts() -> pd.DataFrame:
    series = pd.read_csv(LONGITUDINAL_SERIES_CSV)
    corporate = series.loc[series["series_role"] == "corporate"].copy()
    counts = corporate[
        [
            "macro_topic",
            "corporate_subgroup",
            "corporate_final_merge_group_id",
            "corporate_unique_document_count",
        ]
    ].drop_duplicates()
    counts = counts.rename(
        columns={
            "corporate_subgroup": "subgroup",
            "corporate_final_merge_group_id": "final_merge_group_id",
            "corporate_unique_document_count": "corporate_document_count",
        }
    )
    counts["final_merge_group_id"] = counts["final_merge_group_id"].astype(str)
    return counts


def load_external_aggregate_counts() -> pd.DataFrame:
    series = pd.read_csv(LONGITUDINAL_SERIES_CSV)
    external = series.loc[series["series_role"].isin(["academic_aggregate", "media_aggregate"])].copy()
    if external.empty:
        return pd.DataFrame()
    external["external_source"] = external["series_role"].str.replace("_aggregate", "", regex=False)
    counts = external[
        [
            "macro_topic",
            "corporate_subgroup",
            "corporate_final_merge_group_id",
            "external_source",
            "series_unique_document_count",
            "n_contributing_external_topics",
        ]
    ].drop_duplicates()
    counts = counts.rename(
        columns={
            "corporate_subgroup": "subgroup_corporate",
            "corporate_final_merge_group_id": "final_merge_group_id_corporate",
            "series_unique_document_count": "aggregate_document_count",
            "n_contributing_external_topics": "aggregate_topic_count",
        }
    )
    counts["final_merge_group_id_corporate"] = counts["final_merge_group_id_corporate"].astype(str)
    return counts


def build_corporate_anchors(groups: pd.DataFrame, corporate_counts: pd.DataFrame) -> pd.DataFrame:
    corporate = groups.loc[groups["source"] == "corporate"].copy()
    corporate["final_merge_group_id"] = corporate["final_merge_group_id"].astype(str)
    result = corporate.merge(
        corporate_counts,
        on=["macro_topic", "subgroup", "final_merge_group_id"],
        how="left",
    )
    result["corporate_document_count"] = (
        pd.to_numeric(result["corporate_document_count"], errors="coerce")
        .fillna(pd.to_numeric(result["topic_size"], errors="coerce"))
        .fillna(0)
        .astype(int)
    )
    result["macro_topic_name"] = [
        clean_text(name) or MACRO_TOPIC_NAMES.get(str(code), str(code))
        for code, name in zip(result["macro_topic"], result.get("macro_topic_name", ""))
    ]
    result = result.rename(
        columns={
            "topic_label_refined": "corporate_topic_label",
            "overall_summary": "corporate_summary",
            "topic_name_original": "corporate_topic_name_original",
            "group_size": "merged_microtopic_count",
        }
    )
    output_columns = [
        "macro_topic",
        "macro_topic_name",
        "subgroup",
        "final_merge_group_id",
        "micro_topic_id",
        "corporate_topic_label",
        "corporate_topic_name_original",
        "corporate_summary",
        "corporate_document_count",
        "merged_microtopic_count",
        "active_year_count",
        "active_year_min",
        "active_year_max",
        "source_file",
    ]
    result["source_file"] = f"{source_label(CORPORATE_GROUPS_CSV)}; {source_label(LONGITUDINAL_SERIES_CSV)}"
    return result[output_columns].sort_values(["macro_topic", "corporate_topic_label"], kind="mergesort")


def build_external_topics(aligned: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "macro_topic",
        "macro_topic_name",
        "subgroup_corporate",
        "final_merge_group_id_corporate",
        "micro_topic_id_corporate",
        "topic_label_refined_corporate",
        "overall_summary_corporate",
        "corporate_unique_document_count",
        "source_noncorporate",
        "subgroup_noncorporate",
        "final_merge_group_id_noncorporate",
        "micro_topic_id_noncorporate",
        "topic_label_refined_noncorporate",
        "topic_name_original_noncorporate",
        "overall_summary_noncorporate",
        "external_unique_document_count",
        "external_topic_size_count",
        "external_group_size_count",
        "best_cosine_similarity",
        "direct_pair_count",
        "best_pair_id",
        "inclusion_reason",
        "source_file",
    ]
    result = aligned[columns].copy()
    result = result.rename(
        columns={
            "final_merge_group_id_corporate": "corporate_final_merge_group_id",
            "micro_topic_id_corporate": "corporate_micro_topic_id",
            "topic_label_refined_corporate": "corporate_topic_label",
            "overall_summary_corporate": "corporate_summary",
            "corporate_unique_document_count": "corporate_document_count",
            "source_noncorporate": "external_source",
            "final_merge_group_id_noncorporate": "external_final_merge_group_id",
            "micro_topic_id_noncorporate": "external_micro_topic_id",
            "topic_label_refined_noncorporate": "external_topic_label",
            "topic_name_original_noncorporate": "external_topic_name_original",
            "overall_summary_noncorporate": "external_summary",
            "external_unique_document_count": "external_document_count",
            "external_group_size_count": "external_merged_microtopic_count",
        }
    )
    return result.sort_values(
        ["macro_topic", "corporate_topic_label", "external_source", "external_topic_label"],
        kind="mergesort",
    )


def aggregate_source_rows(aligned: pd.DataFrame, aggregate_counts: pd.DataFrame) -> pd.DataFrame:
    if aligned.empty:
        return pd.DataFrame()
    grouped = (
        aligned.groupby(
            [
                "macro_topic",
                "subgroup_corporate",
                "final_merge_group_id_corporate",
                "source_noncorporate",
            ],
            dropna=False,
        )
        .agg(
            topic_count=("final_merge_group_id_noncorporate", "nunique"),
            summed_external_document_count=("external_unique_document_count", numeric_sum),
            labels=("topic_label_refined_noncorporate", join_values),
            summaries=("overall_summary_noncorporate", join_values),
            source_file=("source_file", join_values),
        )
        .reset_index()
    )
    grouped = grouped.rename(
        columns={
            "source_noncorporate": "external_source",
            "final_merge_group_id_corporate": "final_merge_group_id_corporate",
        }
    )
    if not aggregate_counts.empty:
        grouped = grouped.merge(
            aggregate_counts,
            on=[
                "macro_topic",
                "subgroup_corporate",
                "final_merge_group_id_corporate",
                "external_source",
            ],
            how="left",
        )
    grouped["aggregate_topic_count"] = (
        pd.to_numeric(grouped.get("aggregate_topic_count"), errors="coerce")
        .fillna(pd.to_numeric(grouped["topic_count"], errors="coerce"))
        .fillna(0)
        .astype(int)
    )
    grouped["aggregate_document_count"] = (
        pd.to_numeric(grouped.get("aggregate_document_count"), errors="coerce")
        .fillna(pd.to_numeric(grouped["summed_external_document_count"], errors="coerce"))
        .fillna(0)
        .astype(int)
    )
    grouped["source_file"] = grouped["source_file"].map(
        lambda text: f"{text}; {source_label(LONGITUDINAL_SERIES_CSV)}"
    )
    return grouped


def build_corporate_with_aggregates(
    corporate_anchors: pd.DataFrame,
    aligned: pd.DataFrame,
    aggregate_counts: pd.DataFrame,
) -> pd.DataFrame:
    result = corporate_anchors[
        [
            "macro_topic",
            "macro_topic_name",
            "subgroup",
            "final_merge_group_id",
            "corporate_topic_label",
            "corporate_summary",
            "corporate_document_count",
            "source_file",
        ]
    ].copy()
    result = result.rename(
        columns={
            "subgroup": "subgroup_corporate",
            "final_merge_group_id": "final_merge_group_id_corporate",
            "source_file": "corporate_source_file",
        }
    )
    aggregate_rows = aggregate_source_rows(aligned, aggregate_counts)
    for external_source in ["academic", "media"]:
        source_rows = aggregate_rows.loc[aggregate_rows["external_source"] == external_source].copy()
        source_rows = source_rows.rename(
            columns={
                "aggregate_topic_count": f"{external_source}_topic_count",
                "aggregate_document_count": f"{external_source}_document_count",
                "labels": f"{external_source}_labels",
                "summaries": f"{external_source}_summaries",
                "source_file": f"{external_source}_source_file",
            }
        )
        keep = [
            "macro_topic",
            "subgroup_corporate",
            "final_merge_group_id_corporate",
            f"{external_source}_topic_count",
            f"{external_source}_document_count",
            f"{external_source}_labels",
            f"{external_source}_summaries",
            f"{external_source}_source_file",
        ]
        result = result.merge(
            source_rows[keep],
            on=["macro_topic", "subgroup_corporate", "final_merge_group_id_corporate"],
            how="left",
        )
    for source in ["academic", "media"]:
        result[f"{source}_topic_count"] = pd.to_numeric(
            result[f"{source}_topic_count"], errors="coerce"
        ).fillna(0).astype(int)
        result[f"{source}_document_count"] = pd.to_numeric(
            result[f"{source}_document_count"], errors="coerce"
        ).fillna(0).astype(int)
        result[f"{source}_labels"] = result[f"{source}_labels"].fillna("")
        result[f"{source}_summaries"] = result[f"{source}_summaries"].fillna("")
        result[f"{source}_source_file"] = result[f"{source}_source_file"].fillna("")
    result["source_file"] = result[
        ["corporate_source_file", "academic_source_file", "media_source_file"]
    ].apply(lambda row: join_values(row), axis=1)
    result = result.drop(columns=["corporate_source_file", "academic_source_file", "media_source_file"])
    return result.sort_values(["macro_topic", "corporate_topic_label"], kind="mergesort")


def autosize_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:80])
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 70)
        for row in worksheet.iter_rows():
            for cell in row:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "top"
                cell.alignment = alignment
    workbook.save(path)


def validate_outputs(
    macro_topics: pd.DataFrame,
    macro_expressions: pd.DataFrame,
    corporate_anchors: pd.DataFrame,
    external_topics: pd.DataFrame,
) -> dict[str, int]:
    counts = {
        "macro_topics": int(len(macro_topics)),
        "macro_expressions_long": int(len(macro_expressions)),
        "corporate_anchors": int(len(corporate_anchors)),
        "external_aggregate_topics": int(len(external_topics)),
    }
    if counts["macro_topics"] != 6:
        raise ValueError(f"Expected 6 macro topics, found {counts['macro_topics']}")
    if counts["corporate_anchors"] != 27:
        raise ValueError(f"Expected 27 corporate anchors, found {counts['corporate_anchors']}")
    if counts["external_aggregate_topics"] <= 0:
        raise ValueError("Expected at least one external aligned topic")
    for name, frame in {
        "Macro topics": macro_topics,
        "Macro expressions long": macro_expressions,
        "Corporate anchors": corporate_anchors,
        "External aggregate topics": external_topics,
    }.items():
        if "source_file" not in frame.columns:
            raise ValueError(f"{name} is missing source_file")
    return counts


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    catalog = pd.read_csv(CATALOG_CSV)
    groups = pd.read_csv(CORPORATE_GROUPS_CSV)
    aligned, topic_table_paths = load_aligned_topics()
    corporate_counts = load_corporate_document_counts()
    aggregate_counts = load_external_aggregate_counts()

    readme = build_readme()
    sources = build_sources(topic_table_paths)
    macro_topics = build_macro_topics(catalog)
    macro_expressions = build_macro_expressions(catalog)
    corporate_anchors = build_corporate_anchors(groups, corporate_counts)
    corporate_with_aggregates = build_corporate_with_aggregates(
        corporate_anchors,
        aligned,
        aggregate_counts,
    )
    external_topics = build_external_topics(aligned)
    counts = validate_outputs(macro_topics, macro_expressions, corporate_anchors, external_topics)

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        sources.to_excel(writer, sheet_name="Sources", index=False)
        macro_topics.to_excel(writer, sheet_name="Macro topics", index=False)
        macro_expressions.to_excel(writer, sheet_name="Macro expressions long", index=False)
        corporate_anchors.to_excel(writer, sheet_name="Corporate anchors", index=False)
        corporate_with_aggregates.to_excel(writer, sheet_name="Corporate with aggregates", index=False)
        external_topics.to_excel(writer, sheet_name="External aggregate topics", index=False)
    autosize_workbook(OUTPUT_XLSX)

    manifest = {
        "excel": str(OUTPUT_XLSX),
        "row_counts": counts,
        "source_files": [source_label(CATALOG_CSV), source_label(CORPORATE_GROUPS_CSV), source_label(LONGITUDINAL_SERIES_CSV)]
        + [source_label(path) for path in topic_table_paths],
    }
    (OUTPUT_DIR / "paper_appendix_method_tables_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
