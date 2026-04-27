#!/usr/bin/env python3
"""Build simplified same-year Spearman correlation tables for source-topic relations."""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path

import pandas as pd


PIPELINE_ROOT = Path("paper_pipeline")
RELATION_ROOT = PIPELINE_ROOT / "outputs" / "paper_tables" / "corporate_focus_source_topic_relations"
AGGREGATE_DETAILS_PATH = RELATION_ROOT / "aggregate_source_relations_lag_details.csv"
INDIVIDUAL_DETAILS_PATH = RELATION_ROOT / "individual_source_relations_lag_details.csv"
OUTPUT_XLSX_PATH = RELATION_ROOT / "zero_lag_spearman_relations.xlsx"
OUTPUT_AGGREGATE_CSV_PATH = RELATION_ROOT / "zero_lag_spearman_aggregate_relations.csv"
OUTPUT_INDIVIDUAL_CSV_PATH = RELATION_ROOT / "zero_lag_spearman_individual_relations.csv"


def relation_strength(value: object) -> str:
    converted = pd.to_numeric(value, errors="coerce")
    if pd.isna(converted):
        return "not_available"
    absolute = abs(float(converted))
    if absolute >= 0.70:
        return "strong"
    if absolute >= 0.50:
        return "moderate"
    if absolute >= 0.30:
        return "weak"
    return "very_weak"


def relation_direction(value: object) -> str:
    converted = pd.to_numeric(value, errors="coerce")
    if pd.isna(converted):
        return "not_available"
    if float(converted) > 0:
        return "same_direction"
    if float(converted) < 0:
        return "opposite_direction"
    return "no_monotonic_association"


def significance_label(value: object) -> str:
    converted = pd.to_numeric(value, errors="coerce")
    if pd.isna(converted):
        return "not_significant"
    if float(converted) < 0.05:
        return "q_lt_0_05"
    if float(converted) < 0.10:
        return "q_lt_0_10"
    return "not_significant"


def select_zero_lag(path: Path, level: str) -> pd.DataFrame:
    details = pd.read_csv(path)
    zero_lag = details.loc[details["lag"].astype(int) == 0].copy()
    zero_lag["relation_level"] = level
    zero_lag["same_year_spearman_r"] = pd.to_numeric(zero_lag["spearman_r"], errors="coerce")
    zero_lag["same_year_spearman_p"] = pd.to_numeric(zero_lag["spearman_p"], errors="coerce")
    zero_lag["same_year_spearman_q_table"] = pd.to_numeric(zero_lag["spearman_q_table"], errors="coerce")
    zero_lag["same_year_pearson_r"] = pd.to_numeric(zero_lag["pearson_r"], errors="coerce")
    zero_lag["same_year_pearson_p"] = pd.to_numeric(zero_lag["pearson_p"], errors="coerce")
    zero_lag["same_year_pearson_q_table"] = pd.to_numeric(zero_lag["pearson_q_table"], errors="coerce")
    zero_lag["same_year_relation_direction"] = zero_lag["same_year_spearman_r"].map(relation_direction)
    zero_lag["same_year_relation_strength"] = zero_lag["same_year_spearman_r"].map(relation_strength)
    zero_lag["same_year_significance"] = zero_lag["same_year_spearman_q_table"].map(significance_label)
    zero_lag["same_year_interpretation"] = zero_lag.apply(interpret_row, axis=1)
    return reorder_columns(zero_lag, level)


def interpret_row(row: pd.Series) -> str:
    source = row.get("external_source", "external")
    corporate = row.get("corporate_topic_label", "corporate topic")
    external = row.get("external_topic_label", row.get("external_relation_label", "external counterpart"))
    direction = row.get("same_year_relation_direction", "not_available")
    strength = row.get("same_year_relation_strength", "not_available")
    significance = row.get("same_year_significance", "not_significant")
    r_value = row.get("same_year_spearman_r")
    q_value = row.get("same_year_spearman_q_table")
    r_text = "not available" if pd.isna(r_value) else f"{float(r_value):.3f}"
    q_text = "not available" if pd.isna(q_value) else f"{float(q_value):.3g}"
    if direction == "same_direction":
        movement = "move in the same direction in the same year"
    elif direction == "opposite_direction":
        movement = "move in opposite directions in the same year"
    else:
        movement = "do not show an interpretable same-year monotonic association"
    return (
        f"The corporate topic \"{corporate}\" and the {source} counterpart \"{external}\" {movement}. "
        f"The same-year Spearman association is {strength} (r={r_text}, table-level q={q_text}; {significance})."
    )


def reorder_columns(frame: pd.DataFrame, level: str) -> pd.DataFrame:
    common = [
        "relation_id",
        "relation_level",
        "macro_topic",
        "macro_topic_name",
        "corporate_final_merge_group_id",
        "corporate_topic_label",
        "external_source",
        "external_relation_label",
        "external_topic_count",
        "corporate_unique_document_count",
        "external_unique_document_count",
        "n_overlap_years",
        "low_information_flag",
        "external_series_constant_flag",
        "corporate_series_constant_flag",
        "same_year_spearman_r",
        "same_year_spearman_p",
        "same_year_spearman_q_table",
        "same_year_pearson_r",
        "same_year_pearson_p",
        "same_year_pearson_q_table",
        "same_year_relation_direction",
        "same_year_relation_strength",
        "same_year_significance",
        "same_year_interpretation",
        "corporate_first_active_year",
        "external_first_active_year",
        "corporate_peak_year",
        "external_peak_year",
        "first_active_year_gap",
        "peak_year_gap",
    ]
    individual = [
        "external_subgroup",
        "external_final_merge_group_id",
        "external_micro_topic_id",
        "external_topic_label",
        "best_cosine_similarity",
        "direct_pair_count",
        "best_pair_id",
        "inclusion_reason",
    ]
    columns = common + (individual if level == "individual" else [])
    columns = [column for column in columns if column in frame.columns]
    columns += [column for column in frame.columns if column not in columns]
    return frame[columns]


def autosize_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:80])
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 55)
        for row in worksheet.iter_rows():
            for cell in row:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "top"
                cell.alignment = alignment
    workbook.save(path)


def main() -> None:
    aggregate = select_zero_lag(AGGREGATE_DETAILS_PATH, "aggregate")
    individual = select_zero_lag(INDIVIDUAL_DETAILS_PATH, "individual")
    aggregate.to_csv(OUTPUT_AGGREGATE_CSV_PATH, index=False)
    individual.to_csv(OUTPUT_INDIVIDUAL_CSV_PATH, index=False)

    readme = pd.DataFrame(
        [
            {
                "item": "scope",
                "description": "Same-year Spearman correlations only. No lagged best-lag selection is used in this workbook.",
            },
            {
                "item": "metric",
                "description": "Correlations use annual_document_prevalence: topic documents in a year divided by source-domain documents in that same year.",
            },
            {
                "item": "main columns",
                "description": "Use same_year_spearman_r for association direction/strength and same_year_spearman_q_table for table-level corrected significance.",
            },
            {
                "item": "direction",
                "description": "Positive Spearman means the corporate and external series move in the same direction in the same year; negative means opposite directions.",
            },
            {
                "item": "row counts",
                "description": json.dumps(
                    {
                        "aggregate": int(len(aggregate)),
                        "individual": int(len(individual)),
                        "aggregate_q_lt_0_05": int((aggregate["same_year_spearman_q_table"] < 0.05).sum()),
                        "individual_q_lt_0_05": int((individual["same_year_spearman_q_table"] < 0.05).sum()),
                    },
                    ensure_ascii=False,
                ),
            },
        ]
    )

    with pd.ExcelWriter(OUTPUT_XLSX_PATH, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        aggregate.to_excel(writer, sheet_name="Aggregate same-year", index=False)
        individual.to_excel(writer, sheet_name="Individual same-year", index=False)
    autosize_workbook(OUTPUT_XLSX_PATH)
    print(
        json.dumps(
            {
                "excel": str(OUTPUT_XLSX_PATH),
                "aggregate_csv": str(OUTPUT_AGGREGATE_CSV_PATH),
                "individual_csv": str(OUTPUT_INDIVIDUAL_CSV_PATH),
                "aggregate_rows": int(len(aggregate)),
                "individual_rows": int(len(individual)),
                "aggregate_q_lt_0_05": int((aggregate["same_year_spearman_q_table"] < 0.05).sum()),
                "individual_q_lt_0_05": int((individual["same_year_spearman_q_table"] < 0.05).sum()),
            },
            indent=2,
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
