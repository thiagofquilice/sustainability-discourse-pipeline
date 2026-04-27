#!/usr/bin/env python3
"""Build compact Spearman + peak-result tables for the paper."""

from __future__ import annotations

from copy import copy
from pathlib import Path

import pandas as pd

from analyze_corporate_focus_source_topic_relations import (
    DEFAULT_MERGED_ROOT,
    DEFAULT_SERIES_CSV,
    DEFAULT_TOPIC_TABLES_DIR,
    build_aggregate_relation_rows,
    build_individual_relation_rows,
    load_aligned_topic_tables,
    load_individual_external_series,
)


PIPELINE_ROOT = Path("paper_pipeline")
RELATION_ROOT = PIPELINE_ROOT / "outputs" / "paper_tables" / "corporate_focus_source_topic_relations"
AGGREGATE_INPUT_CSV = RELATION_ROOT / "peak_precedence_aggregate_relations.csv"
INDIVIDUAL_INPUT_CSV = RELATION_ROOT / "peak_precedence_individual_relations.csv"
AGGREGATE_OUTPUT_CSV = RELATION_ROOT / "aggregate_spearman_peak_summary_table.csv"
INDIVIDUAL_OUTPUT_CSV = RELATION_ROOT / "individual_spearman_peak_summary_table.csv"
OUTPUT_XLSX = RELATION_ROOT / "spearman_peak_summary_tables.xlsx"


def format_year(value: object) -> str:
    converted = pd.to_numeric(value, errors="coerce")
    if pd.isna(converted):
        return "NA"
    return str(int(converted))


def format_gap(value: object) -> str:
    converted = pd.to_numeric(value, errors="coerce")
    if pd.isna(converted):
        return "NA"
    return str(int(converted))


def build_peak_result(row: pd.Series) -> str:
    label = str(row.get("peak_precedence_label", "")).strip()
    corporate_peak = format_year(row.get("corporate_peak_year"))
    external_peak = format_year(row.get("external_peak_year"))
    gap = format_gap(row.get("peak_year_gap"))
    if label == "external_peaks_before_corporate":
        result = "External peaks before corporate"
    elif label == "corporate_peaks_before_external":
        result = "Corporate peaks before external"
    elif label == "synchronous_peak":
        result = "Synchronous peak"
    else:
        result = "Insufficient peak data"
    return f"{result} (corporate peak: {corporate_peak}; external peak: {external_peak}; gap: {gap})"


def prevalence_at(series: pd.DataFrame, year: int) -> float:
    selected = series.loc[series["year"].astype(int) == int(year), "annual_document_prevalence"]
    if selected.empty:
        return float("nan")
    return float(pd.to_numeric(selected.iloc[0], errors="coerce"))


def prevalence_window(series: pd.DataFrame, start_year: int, end_year: int) -> pd.Series:
    selected = series.loc[
        (series["year"].astype(int) >= int(start_year)) & (series["year"].astype(int) <= int(end_year)),
        "annual_document_prevalence",
    ]
    return pd.to_numeric(selected, errors="coerce").dropna()


def exposure_growth_metrics(row: dict[str, object]) -> dict[str, object]:
    corporate_first = pd.to_numeric(row.get("corporate_first_active_year"), errors="coerce")
    if pd.isna(corporate_first):
        return {
            "external_prior_exposure_3y_sum": float("nan"),
            "external_prior_exposure_3y_mean": float("nan"),
            "corporate_onset_prevalence": float("nan"),
            "corporate_following_3y_max_prevalence": float("nan"),
            "corporate_later_growth_3y": float("nan"),
            "prior_exposure_later_growth_pattern": "insufficient_data",
        }

    start_year = int(corporate_first)
    external_series = row["_external_series"]
    corporate_series = row["_corporate_series"]
    prior = prevalence_window(external_series, start_year - 3, start_year - 1)
    later = prevalence_window(corporate_series, start_year + 1, start_year + 3)
    onset = prevalence_at(corporate_series, start_year)

    prior_sum = float(prior.sum()) if not prior.empty else 0.0
    prior_mean = float(prior.mean()) if not prior.empty else 0.0
    later_max = float(later.max()) if not later.empty else float("nan")
    growth = later_max - onset if pd.notna(later_max) and pd.notna(onset) else float("nan")

    if prior_sum > 0 and pd.notna(growth) and growth > 0:
        pattern = "external_prior_exposure_with_later_corporate_growth"
    elif prior_sum > 0 and pd.notna(growth) and growth <= 0:
        pattern = "external_prior_exposure_without_later_corporate_growth"
    elif prior_sum <= 0 and pd.notna(growth) and growth > 0:
        pattern = "corporate_growth_without_prior_external_exposure"
    elif pd.isna(growth):
        pattern = "insufficient_later_growth_window"
    else:
        pattern = "no_prior_external_exposure_no_later_corporate_growth"

    return {
        "external_prior_exposure_3y_sum": prior_sum,
        "external_prior_exposure_3y_mean": prior_mean,
        "corporate_onset_prevalence": onset,
        "corporate_following_3y_max_prevalence": later_max,
        "corporate_later_growth_3y": growth,
        "prior_exposure_later_growth_pattern": pattern,
    }


def build_exposure_growth_lookup() -> dict[str, dict[str, object]]:
    series = pd.read_csv(DEFAULT_SERIES_CSV)
    relations = load_aligned_topic_tables(DEFAULT_TOPIC_TABLES_DIR)
    individual_series = load_individual_external_series(relations, DEFAULT_MERGED_ROOT)
    relation_rows = build_aggregate_relation_rows(series) + build_individual_relation_rows(
        series,
        relations,
        individual_series,
    )
    lookup: dict[str, dict[str, object]] = {}
    for relation in relation_rows:
        lookup[str(relation["relation_id"])] = exposure_growth_metrics(relation)
    return lookup


def autosize_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 65)
        for row in worksheet.iter_rows():
            for cell in row:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "top"
                cell.alignment = alignment
    workbook.save(path)


def build_table(
    frame: pd.DataFrame,
    include_individual_columns: bool,
    relation_metrics: dict[str, dict[str, object]],
) -> pd.DataFrame:
    data = {
        "macro_topic": frame["macro_topic"],
        "macro_topic_name": frame["macro_topic_name"],
        "corporate_topic_label": frame["corporate_topic_label"],
        "external_source": frame["external_source"],
        "external_relation_label": frame["external_relation_label"],
    }
    if include_individual_columns:
        data.update(
            {
                "external_topic_label": frame["external_topic_label"],
                "external_final_merge_group_id": frame["external_final_merge_group_id"],
            }
        )
    data.update(
        {
            "corporate_unique_document_count": frame["corporate_unique_document_count"],
            "external_unique_document_count": frame["external_unique_document_count"],
            "spearman_r": pd.to_numeric(frame["same_year_spearman_r"], errors="coerce"),
            "spearman_p": pd.to_numeric(frame["same_year_spearman_p"], errors="coerce"),
            "corporate_first_active_year": frame["corporate_first_active_year"],
            "external_first_active_year": frame["external_first_active_year"],
            "first_active_year_gap": frame["first_active_year_gap"],
            "corporate_peak_year": frame["corporate_peak_year"],
            "external_peak_year": frame["external_peak_year"],
            "peak_year_gap": frame["peak_year_gap"],
            "peak_analysis_result": frame.apply(build_peak_result, axis=1),
        }
    )
    result = pd.DataFrame(data)
    exposure_columns = [
        "external_prior_exposure_3y_sum",
        "external_prior_exposure_3y_mean",
        "corporate_onset_prevalence",
        "corporate_following_3y_max_prevalence",
        "corporate_later_growth_3y",
        "prior_exposure_later_growth_pattern",
    ]
    for column in exposure_columns:
        result[column] = frame["relation_id"].map(
            lambda relation_id: relation_metrics.get(str(relation_id), {}).get(column)
        )
    result = result.sort_values(
        ["macro_topic", "corporate_topic_label", "external_source", "external_relation_label"],
        kind="mergesort",
    ).reset_index(drop=True)
    return result


def build_method_descriptions() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "method": "Same-year Spearman",
                "columns": "spearman_r; spearman_p",
                "description": "Spearman rank correlation between corporate annual_document_prevalence and the corresponding external annual_document_prevalence in the same year. Positive values indicate same-direction co-movement; negative values indicate opposite-direction co-movement.",
            },
            {
                "method": "First active year gap",
                "columns": "corporate_first_active_year; external_first_active_year; first_active_year_gap",
                "description": "Difference between the first year with topic documents in corporate disclosure and the first year with topic documents in the external source. Computed as corporate_first_active_year minus external_first_active_year; positive values mean the external topic appeared earlier.",
            },
            {
                "method": "Peak year gap",
                "columns": "corporate_peak_year; external_peak_year; peak_year_gap; peak_analysis_result",
                "description": "Difference between the year of maximum annual_document_prevalence in corporate disclosure and the year of maximum annual_document_prevalence in the external source. Computed as corporate_peak_year minus external_peak_year; positive values mean the external topic peaked earlier. No plus/minus one-year window is applied in this table.",
            },
            {
                "method": "External prior exposure",
                "columns": "external_prior_exposure_3y_sum; external_prior_exposure_3y_mean",
                "description": "Exposure of the external source before corporate uptake. It is calculated as the sum and mean of external annual_document_prevalence during the three years before the corporate first active year.",
            },
            {
                "method": "Corporate later growth",
                "columns": "corporate_onset_prevalence; corporate_following_3y_max_prevalence; corporate_later_growth_3y",
                "description": "Growth in corporate prevalence after the topic first appears in corporate disclosure. It is calculated as the maximum corporate annual_document_prevalence in the three years after the corporate first active year minus corporate prevalence in the first active year.",
            },
            {
                "method": "Combined exposure-growth pattern",
                "columns": "prior_exposure_later_growth_pattern",
                "description": "Interpretive label combining external prior exposure and corporate later growth. It identifies whether external salience existed before corporate uptake and whether corporate prevalence increased afterward.",
            },
        ]
    )


def main() -> None:
    relation_metrics = build_exposure_growth_lookup()
    aggregate = build_table(
        pd.read_csv(AGGREGATE_INPUT_CSV),
        include_individual_columns=False,
        relation_metrics=relation_metrics,
    )
    individual = build_table(
        pd.read_csv(INDIVIDUAL_INPUT_CSV),
        include_individual_columns=True,
        relation_metrics=relation_metrics,
    )
    methods = build_method_descriptions()

    aggregate.to_csv(AGGREGATE_OUTPUT_CSV, index=False)
    individual.to_csv(INDIVIDUAL_OUTPUT_CSV, index=False)
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        methods.to_excel(writer, sheet_name="Method descriptions", index=False)
        aggregate.to_excel(writer, sheet_name="Aggregate Spearman Peak", index=False)
        individual.to_excel(writer, sheet_name="Individual Spearman Peak", index=False)
    autosize_workbook(OUTPUT_XLSX)
    print(f"Wrote {len(aggregate)} rows to {AGGREGATE_OUTPUT_CSV}")
    print(f"Wrote {len(individual)} rows to {INDIVIDUAL_OUTPUT_CSV}")
    print(f"Wrote {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
