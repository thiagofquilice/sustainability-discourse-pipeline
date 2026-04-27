#!/usr/bin/env python3
"""Build peak-year precedence diagnostics for corporate/external topic relations."""

from __future__ import annotations

import argparse
import json
from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import binomtest


PIPELINE_ROOT = Path("paper_pipeline")
DEFAULT_RELATION_ROOT = PIPELINE_ROOT / "outputs" / "paper_tables" / "corporate_focus_source_topic_relations"
DEFAULT_AGGREGATE_CSV = DEFAULT_RELATION_ROOT / "zero_lag_spearman_aggregate_relations.csv"
DEFAULT_INDIVIDUAL_CSV = DEFAULT_RELATION_ROOT / "zero_lag_spearman_individual_relations.csv"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--aggregate-csv", type=Path, default=DEFAULT_AGGREGATE_CSV)
    parser.add_argument("--individual-csv", type=Path, default=DEFAULT_INDIVIDUAL_CSV)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_RELATION_ROOT)
    return parser.parse_args()


def clean_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def as_number(value: object) -> float:
    converted = pd.to_numeric(value, errors="coerce")
    if pd.isna(converted):
        return np.nan
    return float(converted)


def peak_label(gap: object) -> str:
    gap_value = as_number(gap)
    if pd.isna(gap_value):
        return "insufficient_peak_data"
    if gap_value > 0:
        return "external_peaks_before_corporate"
    if gap_value < 0:
        return "corporate_peaks_before_external"
    return "synchronous_peak"


def peak_window_label(gap: object) -> str:
    gap_value = as_number(gap)
    if pd.isna(gap_value):
        return "insufficient_peak_data"
    if abs(gap_value) <= 1:
        return "near_synchronous_peak_abs_le_1_year"
    if gap_value >= 2:
        return "external_peaks_before_corporate_2plus_years"
    return "corporate_peaks_before_external_2plus_years"


def onset_label(gap: object) -> str:
    gap_value = as_number(gap)
    if pd.isna(gap_value):
        return "insufficient_onset_data"
    if gap_value > 0:
        return "external_active_before_corporate"
    if gap_value < 0:
        return "corporate_active_before_external"
    return "same_first_active_year"


def combined_pattern(row: pd.Series) -> str:
    gap = as_number(row.get("peak_year_gap"))
    rho = as_number(row.get("same_year_spearman_r"))
    if pd.isna(gap) or pd.isna(rho):
        return "insufficient_data"
    if abs(gap) <= 1:
        if rho > 0:
            return "near_synchronous_peak_same_direction"
        if rho < 0:
            return "near_synchronous_peak_opposite_direction"
        return "near_synchronous_peak_no_direction"
    if gap >= 2:
        if rho < 0:
            return "external_prior_peak_inverse_trajectories"
        if rho > 0:
            return "external_prior_peak_same_direction"
        return "external_prior_peak_no_monotonic_direction"
    if rho < 0:
        return "corporate_prior_peak_inverse_trajectories"
    if rho > 0:
        return "corporate_prior_peak_same_direction"
    return "corporate_prior_peak_no_monotonic_direction"


def significance_label(value: object) -> str:
    q_value = as_number(value)
    if pd.isna(q_value):
        return "not_significant"
    if q_value < 0.05:
        return "q_lt_0_05"
    if q_value < 0.10:
        return "q_lt_0_10"
    return "not_significant"


def interpretation(row: pd.Series) -> str:
    corporate = clean_text(row.get("corporate_topic_label")) or "the corporate topic"
    external = clean_text(row.get("external_topic_label")) or clean_text(row.get("external_relation_label")) or "the external counterpart"
    source = clean_text(row.get("external_source")) or "external"
    gap = as_number(row.get("peak_year_gap"))
    rho = as_number(row.get("same_year_spearman_r"))
    q_value = as_number(row.get("same_year_spearman_q_table"))
    if pd.isna(gap):
        return (
            f"Peak timing could not be assessed for \"{corporate}\" and the {source} counterpart "
            f"\"{external}\" because at least one series has no active peak year."
        )

    q_text = "not available" if pd.isna(q_value) else f"{q_value:.3g}"
    rho_text = "not available" if pd.isna(rho) else f"{rho:.3f}"
    if gap > 1:
        timing = (
            f"the {source} counterpart peaks {int(gap)} years before the corporate anchor. "
            "This is consistent with an external issue becoming salient before later corporate disclosure attention"
        )
    elif gap == 1:
        timing = (
            f"the {source} counterpart peaks one year before the corporate anchor. "
            "This indicates a near-synchronous pattern with slight external precedence"
        )
    elif gap == 0:
        timing = "both series peak in the same year, indicating synchronous maximum salience"
    elif gap == -1:
        timing = (
            "the corporate anchor peaks one year before the external counterpart. "
            "This indicates a near-synchronous pattern with slight corporate precedence"
        )
    else:
        timing = (
            f"the corporate anchor peaks {int(abs(gap))} years before the {source} counterpart. "
            "This is consistent with corporate salience preceding the external maximum"
        )

    if pd.isna(rho):
        association = "The same-year Spearman association is not available."
    elif rho < 0:
        association = (
            f"The same-year Spearman association is negative (rho={rho_text}, q={q_text}), "
            "so the peak result should be read as a timing contrast rather than contemporaneous co-movement."
        )
    elif rho > 0:
        association = (
            f"The same-year Spearman association is positive (rho={rho_text}, q={q_text}), "
            "so the peak result complements a pattern of contemporaneous co-movement."
        )
    else:
        association = f"The same-year Spearman association is zero (rho={rho_text}, q={q_text})."

    return f"For \"{corporate}\" and \"{external}\", {timing}. {association}"


def enrich(frame: pd.DataFrame, level: str) -> pd.DataFrame:
    result = frame.copy()
    result["relation_level"] = level
    numeric_columns = [
        "same_year_spearman_r",
        "same_year_spearman_p",
        "same_year_spearman_q_table",
        "corporate_first_active_year",
        "external_first_active_year",
        "corporate_peak_year",
        "external_peak_year",
        "first_active_year_gap",
        "peak_year_gap",
        "corporate_peak_annual_document_prevalence",
        "external_peak_annual_document_prevalence",
        "corporate_unique_document_count",
        "external_unique_document_count",
        "external_topic_count",
    ]
    for column in numeric_columns:
        if column in result.columns:
            result[column] = pd.to_numeric(result[column], errors="coerce")

    if "external_topic_label" not in result.columns:
        result["external_topic_label"] = result.get("external_relation_label", "")

    result["peak_precedence_label"] = result["peak_year_gap"].map(peak_label)
    result["peak_window_label"] = result["peak_year_gap"].map(peak_window_label)
    result["first_active_precedence_label"] = result["first_active_year_gap"].map(onset_label)
    result["same_year_spearman_significance"] = result["same_year_spearman_q_table"].map(significance_label)
    result["peak_gap_abs_years"] = result["peak_year_gap"].abs()
    result["external_peak_precedes_by_2plus_years"] = result["peak_year_gap"] >= 2
    result["corporate_peak_precedes_by_2plus_years"] = result["peak_year_gap"] <= -2
    result["inverse_with_external_prior_peak"] = (
        (result["same_year_spearman_r"] < 0) & (result["peak_year_gap"] >= 2)
    )
    result["combined_peak_correlation_pattern"] = result.apply(combined_pattern, axis=1)
    result["peak_timing_interpretation"] = result.apply(interpretation, axis=1)

    preferred_columns = [
        "relation_id",
        "relation_level",
        "macro_topic",
        "macro_topic_name",
        "corporate_final_merge_group_id",
        "corporate_topic_label",
        "external_source",
        "external_relation_label",
        "external_topic_label",
        "external_topic_count",
        "corporate_unique_document_count",
        "external_unique_document_count",
        "corporate_first_active_year",
        "external_first_active_year",
        "first_active_year_gap",
        "first_active_precedence_label",
        "corporate_peak_year",
        "external_peak_year",
        "peak_year_gap",
        "peak_gap_abs_years",
        "peak_precedence_label",
        "peak_window_label",
        "corporate_peak_annual_document_prevalence",
        "external_peak_annual_document_prevalence",
        "same_year_spearman_r",
        "same_year_spearman_p",
        "same_year_spearman_q_table",
        "same_year_spearman_significance",
        "combined_peak_correlation_pattern",
        "external_peak_precedes_by_2plus_years",
        "corporate_peak_precedes_by_2plus_years",
        "inverse_with_external_prior_peak",
        "peak_timing_interpretation",
        "external_subgroup",
        "external_final_merge_group_id",
        "external_micro_topic_id",
        "best_cosine_similarity",
        "inclusion_reason",
    ]
    ordered = [column for column in preferred_columns if column in result.columns]
    ordered += [column for column in result.columns if column not in ordered]
    return result[ordered]


def binomial_p(k: int, n: int, alternative: str) -> float:
    if n <= 0:
        return np.nan
    return float(binomtest(k, n, p=0.5, alternative=alternative).pvalue)


def summarize_subset(frame: pd.DataFrame, label: str) -> dict[str, object]:
    valid = frame.loc[frame["peak_precedence_label"] != "insufficient_peak_data"].copy()
    external_strict = int((valid["peak_year_gap"] > 0).sum())
    corporate_strict = int((valid["peak_year_gap"] < 0).sum())
    synchronous = int((valid["peak_year_gap"] == 0).sum())
    external_window = int((valid["peak_year_gap"] >= 2).sum())
    corporate_window = int((valid["peak_year_gap"] <= -2).sum())
    near_sync = int((valid["peak_year_gap"].abs() <= 1).sum())
    strict_n = external_strict + corporate_strict
    window_n = external_window + corporate_window
    return {
        "comparison_set": label,
        "rows": int(len(frame)),
        "valid_peak_rows": int(len(valid)),
        "external_peaks_before_corporate": external_strict,
        "corporate_peaks_before_external": corporate_strict,
        "synchronous_peak": synchronous,
        "external_peaks_before_corporate_2plus_years": external_window,
        "corporate_peaks_before_external_2plus_years": corporate_window,
        "near_synchronous_peak_abs_le_1_year": near_sync,
        "strict_directional_rows": strict_n,
        "strict_external_share": external_strict / strict_n if strict_n else np.nan,
        "strict_binomial_p_two_sided": binomial_p(external_strict, strict_n, "two-sided"),
        "strict_binomial_p_external_more_frequent": binomial_p(external_strict, strict_n, "greater"),
        "strict_binomial_p_corporate_more_frequent": binomial_p(external_strict, strict_n, "less"),
        "window_directional_rows": window_n,
        "window_external_share": external_window / window_n if window_n else np.nan,
        "window_binomial_p_two_sided": binomial_p(external_window, window_n, "two-sided"),
        "window_binomial_p_external_more_frequent": binomial_p(external_window, window_n, "greater"),
        "window_binomial_p_corporate_more_frequent": binomial_p(external_window, window_n, "less"),
        "inverse_with_external_prior_peak": int(valid["inverse_with_external_prior_peak"].sum()),
    }


def build_summary(aggregate: pd.DataFrame, individual: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for level, frame in [("aggregate", aggregate), ("individual", individual)]:
        rows.append(summarize_subset(frame, f"{level}: all"))
        for source, source_frame in frame.groupby("external_source", dropna=False):
            rows.append(summarize_subset(source_frame, f"{level}: {source}"))
    return pd.DataFrame(rows)


def autosize_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:80])
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 70)
        for row in worksheet.iter_rows():
            for cell in row:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "top"
                cell.alignment = alignment
    workbook.save(path)


def build_readme(aggregate: pd.DataFrame, individual: pd.DataFrame, summary: pd.DataFrame) -> pd.DataFrame:
    payload = {
        "aggregate_rows": int(len(aggregate)),
        "individual_rows": int(len(individual)),
        "aggregate_external_peak_2plus": int((aggregate["peak_year_gap"] >= 2).sum()),
        "individual_external_peak_2plus": int((individual["peak_year_gap"] >= 2).sum()),
        "aggregate_inverse_external_prior_peak": int(aggregate["inverse_with_external_prior_peak"].sum()),
        "individual_inverse_external_prior_peak": int(individual["inverse_with_external_prior_peak"].sum()),
    }
    return pd.DataFrame(
        [
            {
                "item": "scope",
                "description": "Peak-year timing diagnostic for corporate/external topic relations using annual_document_prevalence.",
            },
            {
                "item": "peak_year_gap",
                "description": "corporate_peak_year minus external_peak_year. Positive values mean the external source peaked before corporate; negative values mean corporate peaked before the external source.",
            },
            {
                "item": "windowed classification",
                "description": "The workbook treats absolute gaps of 0 or 1 year as near-synchronous. Gaps of at least 2 years are labeled as clearer external or corporate precedence.",
            },
            {
                "item": "binomial summary",
                "description": "The Summary sheet reports simple binomial sign tests asking whether external-before peaks are more frequent than corporate-before peaks within each comparison set. This is a descriptive timing test, not causal identification.",
            },
            {
                "item": "correlation combination",
                "description": "combined_peak_correlation_pattern combines peak timing with same-year Spearman. inverse_with_external_prior_peak identifies cases where external peaks earlier but the same-year association is negative.",
            },
            {
                "item": "row counts",
                "description": json.dumps(payload, ensure_ascii=False),
            },
        ]
    )


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    aggregate = enrich(pd.read_csv(args.aggregate_csv), "aggregate")
    individual = enrich(pd.read_csv(args.individual_csv), "individual")
    summary = build_summary(aggregate, individual)
    readme = build_readme(aggregate, individual, summary)

    aggregate_csv = args.output_dir / "peak_precedence_aggregate_relations.csv"
    individual_csv = args.output_dir / "peak_precedence_individual_relations.csv"
    xlsx_path = args.output_dir / "peak_precedence_relations.xlsx"

    aggregate.to_csv(aggregate_csv, index=False)
    individual.to_csv(individual_csv, index=False)

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)
        aggregate.to_excel(writer, sheet_name="Aggregate peak test", index=False)
        individual.to_excel(writer, sheet_name="Individual peak test", index=False)
    autosize_workbook(xlsx_path)

    print(
        json.dumps(
            {
                "excel": str(xlsx_path),
                "aggregate_csv": str(aggregate_csv),
                "individual_csv": str(individual_csv),
                "aggregate_rows": int(len(aggregate)),
                "individual_rows": int(len(individual)),
                "summary": summary.to_dict(orient="records"),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
