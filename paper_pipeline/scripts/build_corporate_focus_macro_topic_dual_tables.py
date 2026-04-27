#!/usr/bin/env python3
"""Build two paper-facing tables per macro topic for the corporate-focus package."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font


PIPELINE_ROOT = Path("paper_pipeline")
REVIEW_ROOT = PIPELINE_ROOT / "outputs" / "corporate_focus_review_with_overrides"
MERGED_ROOT = PIPELINE_ROOT / "outputs" / "bertopic_micro_merged_multiaspect_reviewed"
OUTPUT_DIR = PIPELINE_ROOT / "outputs" / "paper_tables" / "corporate_focus_macro_topic_dual_tables"

ALL_GROUPS_PATH = REVIEW_ROOT / "corporate_focus_all_groups_optional.csv"
INCLUDED_CORPORATE_PATH = REVIEW_ROOT / "included_corporate_groups.csv"
MASTER_REVIEW_PATH = REVIEW_ROOT / "corporate_focus_master_review.csv"
COMMENTED_DECISIONS_PATH = REVIEW_ROOT / "corporate_focus_commented_decision_rows.csv"

WORKBOOK_PATH = OUTPUT_DIR / "corporate_focus_macro_topic_dual_tables.xlsx"
MANIFEST_PATH = OUTPUT_DIR / "corporate_focus_macro_topic_dual_tables_manifest.json"

MACRO_TOPIC_ORDER = ["T1", "T2", "T3", "T4", "T5", "T6"]
SOURCE_ORDER = ["academic", "media"]
MACRO_TOPIC_NAME_MAP = {
    "T1": "Clean energy transition",
    "T2": "Operational sustainability and circular production",
    "T3": "Sustainable products, services and consumption",
    "T4": "Climate strategy, carbon governance and disclosure",
    "T5": "Climate risk, adaptation and resilience",
    "T6": "Ecosystems, pollution and environmental stewardship",
}

CORPORATE_ROSTER_COLUMNS = [
    "macro_topic",
    "macro_topic_name",
    "corporate_topic_label",
    "corporate_summary",
    "topic_size_count",
    "unique_document_count",
    "group_size",
    "active_year_min",
    "active_year_max",
    "subgroup",
    "final_merge_group_id",
]

EXTERNAL_COLUMNS = [
    "macro_topic",
    "macro_topic_name",
    "table_section",
    "corporate_topic_label",
    "source",
    "external_topic_label",
    "external_summary",
    "topic_size_count",
    "unique_document_count",
    "best_cosine_similarity",
    "precedence_type_label",
    "inclusion_reason",
    "subgroup_noncorporate",
    "final_merge_group_id_noncorporate",
    "final_merge_group_id_corporate",
]


def ensure_directory(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def normalize_review_decision(value: object) -> str:
    if not isinstance(value, str):
        return ""
    text = value.strip().lower()
    if text == "delete":
        return "delete"
    if text == "not related":
        return "not related"
    return ""


def detect_document_id_column(frame: pd.DataFrame) -> str:
    for name in ["source_doc_id", "document_id", "doc_id"]:
        if name in frame.columns:
            return name
    raise KeyError("Could not find a stable document identifier column in document_topics.csv")


def load_group_document_counts(all_groups: pd.DataFrame) -> tuple[dict[tuple[str, str], int], str]:
    relevant = all_groups[["subgroup", "final_merge_group_id"]].drop_duplicates().copy()
    relevant["final_merge_group_id"] = relevant["final_merge_group_id"].astype(str)
    group_key_set = set(
        tuple(row)
        for row in relevant[["subgroup", "final_merge_group_id"]].itertuples(index=False, name=None)
    )

    counts: dict[tuple[str, str], int] = {}
    document_id_column_used = ""

    for subgroup in sorted(relevant["subgroup"].unique()):
        document_topics_path = MERGED_ROOT / subgroup / "document_topics.csv"
        document_topics = pd.read_csv(document_topics_path)
        doc_col = detect_document_id_column(document_topics)
        document_id_column_used = doc_col
        document_topics = document_topics.rename(columns={doc_col: "stable_doc_id"})
        document_topics["final_merge_group_id"] = document_topics["final_merge_group_id"].astype(str)
        document_topics["subgroup"] = subgroup

        allowed = {
            key[1]
            for key in group_key_set
            if key[0] == subgroup
        }
        filtered = document_topics.loc[
            document_topics["final_merge_group_id"].isin(allowed)
        ].copy()

        for group_id, group in filtered.groupby("final_merge_group_id", sort=False):
            counts[(subgroup, group_id)] = int(group["stable_doc_id"].dropna().astype(str).nunique())

    return counts, document_id_column_used


def load_all_groups() -> pd.DataFrame:
    frame = pd.read_csv(ALL_GROUPS_PATH)
    frame["macro_topic_name"] = frame["macro_topic"].map(MACRO_TOPIC_NAME_MAP)
    frame["final_merge_group_id"] = frame["final_merge_group_id"].astype(str)
    return frame


def load_corporate_roster(all_groups: pd.DataFrame, group_doc_counts: dict[tuple[str, str], int]) -> pd.DataFrame:
    included_corporate = pd.read_csv(INCLUDED_CORPORATE_PATH)
    included_keys = set(
        tuple(row)
        for row in included_corporate[["subgroup", "final_merge_group_id"]]
        .astype({"final_merge_group_id": str})
        .itertuples(index=False, name=None)
    )

    corporate = all_groups.loc[all_groups["source"] == "corporate"].copy()
    corporate = corporate.loc[
        corporate.apply(lambda row: (row["subgroup"], row["final_merge_group_id"]) in included_keys, axis=1)
    ].copy()

    corporate["corporate_topic_label"] = corporate["topic_label_refined"]
    corporate["corporate_summary"] = corporate["overall_summary"]
    corporate["topic_size_count"] = corporate["topic_size"]
    corporate["unique_document_count"] = corporate.apply(
        lambda row: group_doc_counts.get((row["subgroup"], row["final_merge_group_id"]), 0),
        axis=1,
    )

    corporate = corporate[[
        "macro_topic",
        "macro_topic_name",
        "corporate_topic_label",
        "corporate_summary",
        "topic_size_count",
        "unique_document_count",
        "group_size",
        "active_year_min",
        "active_year_max",
        "subgroup",
        "final_merge_group_id",
    ]].copy()

    corporate = corporate.sort_values(
        by=["macro_topic", "topic_size_count", "corporate_topic_label"],
        ascending=[True, False, True],
        kind="mergesort",
    ).reset_index(drop=True)
    return corporate


def load_final_external_status_map() -> dict[str, str]:
    decisions = pd.read_csv(COMMENTED_DECISIONS_PATH)
    decision_column = "_normalized_review_decision" if "_normalized_review_decision" in decisions.columns else "review_decision"
    decisions["_normalized_review_decision"] = decisions[decision_column].map(normalize_review_decision)
    decisions = decisions.loc[decisions["_normalized_review_decision"] != ""].copy()
    if "best_pair_id" not in decisions.columns:
        raise KeyError("Expected best_pair_id in commented decisions file.")
    return dict(
        decisions[["best_pair_id", "_normalized_review_decision"]]
        .drop_duplicates(subset=["best_pair_id"], keep="last")
        .itertuples(index=False, name=None)
    )


def build_external_table(
    all_groups: pd.DataFrame,
    group_doc_counts: dict[tuple[str, str], int],
) -> pd.DataFrame:
    master = pd.read_csv(MASTER_REVIEW_PATH)
    master["macro_topic_name"] = master["macro_topic"].map(MACRO_TOPIC_NAME_MAP)
    decision_map = load_final_external_status_map()
    master["final_status"] = master["best_pair_id"].map(decision_map).fillna("aligned")
    master = master.loc[master["final_status"] != "delete"].copy()

    group_metadata = all_groups[[
        "subgroup",
        "final_merge_group_id",
        "topic_size",
    ]].drop_duplicates().copy()
    group_metadata["final_merge_group_id"] = group_metadata["final_merge_group_id"].astype(str)

    master["final_merge_group_id_noncorporate"] = master["final_merge_group_id_noncorporate"].astype(str)
    master["final_merge_group_id_corporate"] = master["final_merge_group_id_corporate"].astype(str)

    master = master.merge(
        group_metadata.rename(
            columns={
                "subgroup": "subgroup_noncorporate",
                "final_merge_group_id": "final_merge_group_id_noncorporate",
                "topic_size": "topic_size_count",
            }
        ),
        on=["subgroup_noncorporate", "final_merge_group_id_noncorporate"],
        how="left",
        validate="many_to_one",
    )

    master["unique_document_count"] = master.apply(
        lambda row: group_doc_counts.get((row["subgroup_noncorporate"], row["final_merge_group_id_noncorporate"]), 0),
        axis=1,
    )

    master["source"] = master["source_noncorporate"]
    master["external_topic_label"] = master["topic_label_refined_noncorporate"]
    master["external_summary"] = master["overall_summary_noncorporate"]
    master["corporate_topic_label"] = master["topic_label_refined_corporate"]
    master["table_section"] = master["final_status"].map(
        {
            "aligned": "aligned_to_corporate_anchor",
            "not related": "relevant_external_without_pair",
        }
    )
    master.loc[master["table_section"] == "relevant_external_without_pair", "corporate_topic_label"] = ""

    external = master[[
        "macro_topic",
        "macro_topic_name",
        "table_section",
        "corporate_topic_label",
        "source",
        "external_topic_label",
        "external_summary",
        "topic_size_count",
        "unique_document_count",
        "best_cosine_similarity",
        "precedence_type_label",
        "inclusion_reason",
        "subgroup_noncorporate",
        "final_merge_group_id_noncorporate",
        "final_merge_group_id_corporate",
    ]].copy()

    external["source"] = pd.Categorical(external["source"], categories=SOURCE_ORDER, ordered=True)
    aligned = external.loc[external["table_section"] == "aligned_to_corporate_anchor"].copy()
    aligned = aligned.sort_values(
        by=["macro_topic", "corporate_topic_label", "source", "best_cosine_similarity", "external_topic_label"],
        ascending=[True, True, True, False, True],
        kind="mergesort",
    )
    unpaired = external.loc[external["table_section"] == "relevant_external_without_pair"].copy()
    unpaired = unpaired.sort_values(
        by=["macro_topic", "source", "topic_size_count", "external_topic_label"],
        ascending=[True, True, False, True],
        kind="mergesort",
    )
    combined = pd.concat([aligned, unpaired], ignore_index=True)
    combined["source"] = combined["source"].astype(str)
    return combined


def write_excel_workbook(
    corporate_roster: pd.DataFrame,
    external_table: pd.DataFrame,
) -> list[str]:
    written_sheets: list[str] = []
    with pd.ExcelWriter(WORKBOOK_PATH, engine="openpyxl") as writer:
        for macro_topic in MACRO_TOPIC_ORDER:
            roster_sheet = corporate_roster.loc[corporate_roster["macro_topic"] == macro_topic, CORPORATE_ROSTER_COLUMNS].copy()
            external_sheet = external_table.loc[external_table["macro_topic"] == macro_topic, EXTERNAL_COLUMNS].copy()

            sheet_name = macro_topic
            roster_title_row = 0
            roster_startrow = 1
            external_title_row = roster_startrow + len(roster_sheet) + 3
            external_startrow = external_title_row + 1

            roster_sheet.to_excel(writer, sheet_name=sheet_name, startrow=roster_startrow, index=False)
            external_sheet.to_excel(writer, sheet_name=sheet_name, startrow=external_startrow, index=False)

            ws = writer.sheets[sheet_name]
            ws.cell(row=roster_title_row + 1, column=1, value=f"{macro_topic} — Corporate roster")
            ws.cell(row=external_title_row + 1, column=1, value=f"{macro_topic} — External topics by corporate anchor")
            ws.cell(row=roster_title_row + 1, column=1).font = Font(bold=True)
            ws.cell(row=external_title_row + 1, column=1).font = Font(bold=True)
            ws.freeze_panes = "A2"

            max_widths: dict[int, int] = {}
            for frame, start_row in [(roster_sheet, roster_startrow), (external_sheet, external_startrow)]:
                for idx, column in enumerate(frame.columns, start=1):
                    values = [str(column)]
                    values.extend(frame[column].fillna("").astype(str).head(100).tolist())
                    width = min(max(len(value) for value in values) + 2, 60)
                    max_widths[idx] = max(max_widths.get(idx, 0), width)
            for idx, width in max_widths.items():
                ws.column_dimensions[chr(64 + idx)].width = width

            written_sheets.append(sheet_name)
    return written_sheets


def main() -> None:
    ensure_directory(OUTPUT_DIR)

    all_groups = load_all_groups()
    group_doc_counts, document_id_column_used = load_group_document_counts(all_groups)
    corporate_roster = load_corporate_roster(all_groups, group_doc_counts)
    external_table = build_external_table(all_groups, group_doc_counts)

    manifest_rows: dict[str, dict[str, int]] = {}
    roster_expected_counts = (
        pd.read_csv(INCLUDED_CORPORATE_PATH)
        .groupby("assigned_label")
        .size()
        .reindex(MACRO_TOPIC_ORDER, fill_value=0)
        .astype(int)
        .to_dict()
    )

    csv_paths: list[str] = []
    for macro_topic in MACRO_TOPIC_ORDER:
        roster_frame = corporate_roster.loc[
            corporate_roster["macro_topic"] == macro_topic, CORPORATE_ROSTER_COLUMNS
        ].copy()
        external_frame = external_table.loc[
            external_table["macro_topic"] == macro_topic, EXTERNAL_COLUMNS
        ].copy()

        roster_path = OUTPUT_DIR / f"{macro_topic}_corporate_roster.csv"
        external_path = OUTPUT_DIR / f"{macro_topic}_external_by_corporate_anchor.csv"
        roster_frame.to_csv(roster_path, index=False)
        external_frame.to_csv(external_path, index=False)
        csv_paths.extend([str(roster_path), str(external_path)])

        manifest_rows[macro_topic] = {
            "corporate_roster_rows": int(len(roster_frame)),
            "aligned_rows": int((external_frame["table_section"] == "aligned_to_corporate_anchor").sum()),
            "unpaired_rows": int((external_frame["table_section"] == "relevant_external_without_pair").sum()),
            "external_total_rows": int(len(external_frame)),
        }

    written_sheets = write_excel_workbook(corporate_roster, external_table)

    validations = {
        "corporate_roster_counts_match_included_corporate": all(
            manifest_rows[macro]["corporate_roster_rows"] == roster_expected_counts.get(macro, 0)
            for macro in MACRO_TOPIC_ORDER
        ),
        "no_delete_rows_in_external_table": bool(
            (external_table["table_section"].isin(["aligned_to_corporate_anchor", "relevant_external_without_pair"])).all()
        ),
        "all_unique_document_counts_present": bool(
            (corporate_roster["unique_document_count"].notna().all())
            and (external_table["unique_document_count"].notna().all())
        ),
        "workbook_has_six_sheets": len(written_sheets) == 6,
        "csv_count_is_twelve": len(csv_paths) == 12,
        "document_id_column_used": document_id_column_used,
    }

    manifest = {
        "output_dir": str(OUTPUT_DIR),
        "workbook_path": str(WORKBOOK_PATH),
        "source_files": {
            "all_groups": str(ALL_GROUPS_PATH),
            "included_corporate_groups": str(INCLUDED_CORPORATE_PATH),
            "master_review": str(MASTER_REVIEW_PATH),
            "commented_decisions": str(COMMENTED_DECISIONS_PATH),
            "merged_root": str(MERGED_ROOT),
        },
        "csv_paths": csv_paths,
        "macro_topic_row_counts": manifest_rows,
        "validations": validations,
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    print("CORPORATE_ROSTER_COUNTS")
    print(corporate_roster.groupby("macro_topic").size().reindex(MACRO_TOPIC_ORDER, fill_value=0).to_string())
    print("\nEXTERNAL_TABLE_COUNTS")
    print(external_table.groupby(["macro_topic", "table_section"]).size().to_string())
    print("\nVALIDATIONS")
    print(json.dumps(validations, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
