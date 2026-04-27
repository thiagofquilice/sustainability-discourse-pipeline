#!/usr/bin/env python3
"""Prepare and chunk the unified sustainability corpus."""

from __future__ import annotations

import argparse
import logging
import random
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from transformers import AutoTokenizer

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from utils.chunking import Chunk, token_chunk_from_ids
from utils.load_data import locate_input_files, iter_jsonl
from utils.sector_mapping import ALLOWED_SIC2, compute_sic2, map_industry
from utils.text_cleaning import clean_plain_text, extract_html_paragraphs, join_paragraphs


LOGGER = logging.getLogger("prepare_corpus")
TOKENIZER_REPO_ID = "BAAI/bge-large-en-v1.5"
LOCAL_TOKENIZER_DIR = PROJECT_ROOT / ".hf-cache" / "BAAI_bge-large-en-v1.5"
OUTPUT_CORPUS_REL = Path("data/processed/full_corpus_chunked.parquet")
CHUNKING_REPORT_REL = Path("data/processed/chunking_report.md")
VALIDATION_WORKBOOK_REL = Path("data/validation/chunking_validation.xlsx")


@dataclass
class SourceMetrics:
    raw_records: int = 0
    filtered_year: int = 0
    filtered_empty: int = 0
    filtered_sic: int = 0
    kept_documents: int = 0
    chunked_documents: int = 0
    final_chunks: int = 0
    long_documents: int = 0
    paragraph_overflow_documents: int = 0


@dataclass
class SampleState:
    reservoir: list[dict[str, Any]] = field(default_factory=list)
    seen: int = 0


FINAL_SCHEMA = pa.schema(
    [
        ("doc_id", pa.string()),
        ("chunk_id", pa.string()),
        ("text", pa.string()),
        ("year", pa.int32()),
        ("source", pa.string()),
        ("company", pa.string()),
        ("ticker", pa.string()),
        ("sic", pa.string()),
        ("sic2", pa.string()),
        ("industry", pa.string()),
        ("item", pa.string()),
        ("token_count", pa.int32()),
        ("chunk_start_token", pa.int32()),
        ("chunk_end_token", pa.int32()),
        ("source_doc_id", pa.string()),
        ("raw_id", pa.string()),
        ("original_file", pa.string()),
        ("record_position", pa.int64()),
        ("original_token_count", pa.int32()),
        ("chunk_count", pa.int32()),
        ("was_chunked", pa.bool_()),
        ("filing_date", pa.string()),
        ("form_type", pa.string()),
    ]
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--tokenizer-path", type=Path, default=LOCAL_TOKENIZER_DIR)
    parser.add_argument("--max-tokens", type=int, default=350)
    parser.add_argument("--overlap-tokens", type=int, default=50)
    parser.add_argument("--batch-size", type=int, default=128)
    parser.add_argument("--parquet-batch-size", type=int, default=5000)
    parser.add_argument("--validation-sample-size", type=int, default=50)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--log-level", type=str, default="INFO")
    return parser.parse_args()


def load_tokenizer(tokenizer_path: Path):
    if tokenizer_path.exists():
        tokenizer = AutoTokenizer.from_pretrained(tokenizer_path, local_files_only=True, use_fast=True)
        origin = str(tokenizer_path)
    else:
        tokenizer = AutoTokenizer.from_pretrained(TOKENIZER_REPO_ID, use_fast=True)
        origin = TOKENIZER_REPO_ID

    tokenizer.model_max_length = 10**9
    return tokenizer, origin


class ParquetBufferWriter:
    def __init__(self, output_path: Path, schema: pa.Schema, batch_size: int) -> None:
        self.output_path = output_path
        self.temp_path = output_path.with_suffix(output_path.suffix + ".tmp")
        self.schema = schema
        self.batch_size = batch_size
        self.writer: pq.ParquetWriter | None = None
        self.buffer: list[dict[str, Any]] = []

    def append(self, row: dict[str, Any]) -> None:
        self.buffer.append(row)
        if len(self.buffer) >= self.batch_size:
            self.flush()

    def flush(self) -> None:
        if not self.buffer:
            return
        table = pa.Table.from_pylist(self.buffer, schema=self.schema)
        if self.writer is None:
            self.temp_path.parent.mkdir(parents=True, exist_ok=True)
            if self.temp_path.exists():
                self.temp_path.unlink()
            self.writer = pq.ParquetWriter(self.temp_path, self.schema, compression="snappy")
        self.writer.write_table(table)
        self.buffer.clear()

    def close(self) -> None:
        self.flush()
        if self.writer is None:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            empty_table = pa.Table.from_pylist([], schema=self.schema)
            pq.write_table(empty_table, self.output_path, compression="snappy")
            return
        if self.writer is not None:
            self.writer.close()
        self.temp_path.replace(self.output_path)


def normalize_nullable_string(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def reservoir_update(
    sample_state: SampleState,
    candidate: dict[str, Any],
    sample_size: int,
    rng: random.Random,
) -> None:
    sample_state.seen += 1
    if len(sample_state.reservoir) < sample_size:
        sample_state.reservoir.append(candidate)
        return
    replacement_index = rng.randrange(sample_state.seen)
    if replacement_index < sample_size:
        sample_state.reservoir[replacement_index] = candidate


def make_chunk_id(doc_id: str, chunk_number: int) -> str:
    return f"{doc_id}::chunk_{chunk_number:04d}"


def build_final_row(
    base_record: dict[str, Any],
    chunk: Chunk,
    chunk_number: int,
    chunk_count: int,
) -> dict[str, Any]:
    return {
        "doc_id": base_record["doc_id"],
        "chunk_id": make_chunk_id(base_record["doc_id"], chunk_number),
        "text": chunk.text,
        "year": int(base_record["year"]),
        "source": base_record["source"],
        "company": base_record["company"],
        "ticker": base_record["ticker"],
        "sic": base_record["sic"],
        "sic2": base_record["sic2"],
        "industry": base_record["industry"],
        "item": base_record["item"],
        "token_count": int(chunk.token_count),
        "chunk_start_token": int(chunk.chunk_start_token),
        "chunk_end_token": int(chunk.chunk_end_token),
        "source_doc_id": base_record["source_doc_id"],
        "raw_id": base_record["raw_id"],
        "original_file": base_record["original_file"],
        "record_position": int(base_record["record_position"]),
        "original_token_count": int(base_record["original_token_count"]),
        "chunk_count": int(chunk_count),
        "was_chunked": bool(chunk_count > 1),
        "filing_date": base_record["filing_date"],
        "form_type": base_record["form_type"],
    }


def build_validation_candidate(
    final_row: dict[str, Any],
    base_record: dict[str, Any],
) -> dict[str, Any]:
    return {
        "source": final_row["source"],
        "doc_id": final_row["doc_id"],
        "chunk_id": final_row["chunk_id"],
        "year": final_row["year"],
        "company": final_row["company"],
        "sic": final_row["sic"],
        "sic2": final_row["sic2"],
        "industry": final_row["industry"],
        "item": final_row["item"],
        "token_count": final_row["token_count"],
        "chunk_start_token": final_row["chunk_start_token"],
        "chunk_end_token": final_row["chunk_end_token"],
        "original_text_preview": base_record["original_text_preview"],
        "chunk_text": final_row["text"],
        "validation_status": "",
        "validation_ok": "",
        "reviewer_notes": "",
    }


def emit_document(
    base_record: dict[str, Any],
    chunks: list[Chunk],
    writer: ParquetBufferWriter,
    metrics: SourceMetrics,
    sample_state: SampleState,
    sample_size: int,
    rng: random.Random,
) -> None:
    if not chunks:
        return

    metrics.kept_documents += 1
    metrics.final_chunks += len(chunks)
    if len(chunks) > 1:
        metrics.chunked_documents += 1

    for chunk_number, chunk in enumerate(chunks, start=1):
        final_row = build_final_row(base_record, chunk, chunk_number, len(chunks))
        writer.append(final_row)
        if len(chunks) > 1:
            candidate = build_validation_candidate(final_row, base_record)
            reservoir_update(sample_state, candidate, sample_size, rng)


def build_standard_record(
    source: str,
    raw_row: dict[str, Any],
    path: Path,
    line_number: int,
) -> dict[str, Any]:
    if source == "media":
        raw_id = str(raw_row["id"])
        return {
            "doc_id": f"guardian::{raw_id}::row_{line_number}",
            "source_doc_id": raw_id,
            "raw_id": raw_id,
            "source": "media",
            "year": int(raw_row["year"]),
            "company": None,
            "ticker": None,
            "sic": None,
            "sic2": None,
            "industry": None,
            "item": None,
            "text": clean_plain_text(raw_row.get("document")),
            "original_file": path.name,
            "record_position": line_number,
            "filing_date": None,
            "form_type": None,
        }

    if source == "academic":
        raw_id = str(raw_row["id"])
        return {
            "doc_id": f"paper::{raw_id}",
            "source_doc_id": raw_id,
            "raw_id": raw_id,
            "source": "academic",
            "year": int(raw_row["year"]),
            "company": None,
            "ticker": None,
            "sic": None,
            "sic2": None,
            "industry": None,
            "item": None,
            "text": clean_plain_text(raw_row.get("document")),
            "original_file": path.name,
            "record_position": line_number,
            "filing_date": None,
            "form_type": None,
        }

    raw_id = str(raw_row["id"])
    sic_value = raw_row.get("sic_primary")
    sic = str(sic_value) if sic_value is not None else None
    sic2 = compute_sic2(sic_value)
    return {
        "doc_id": f"10k::{raw_id}",
        "source_doc_id": raw_id,
        "raw_id": raw_id,
        "source": "corporate",
        "year": int(raw_row["year"]),
        "company": normalize_nullable_string(raw_row.get("company")),
        "ticker": normalize_nullable_string(raw_row.get("ticker")),
        "sic": sic,
        "sic2": sic2,
        "industry": map_industry(sic2),
        "item": normalize_nullable_string(raw_row.get("item")),
        "text": raw_row.get("content") or "",
        "original_file": path.name,
        "record_position": line_number,
        "filing_date": normalize_nullable_string(raw_row.get("webPublicationDate")),
        "form_type": normalize_nullable_string(raw_row.get("formType")),
    }


def create_single_chunk(text: str, content_token_count: int, special_tokens: int) -> Chunk:
    return Chunk(
        text=text,
        token_count=content_token_count + special_tokens,
        chunk_start_token=0,
        chunk_end_token=content_token_count - 1,
    )


def flush_non_corporate_batch(
    batch_records: list[dict[str, Any]],
    batch_texts: list[str],
    tokenizer,
    special_tokens: int,
    max_tokens: int,
    overlap_tokens: int,
    writer: ParquetBufferWriter,
    metrics: SourceMetrics,
    sample_state: SampleState,
    sample_size: int,
    rng: random.Random,
) -> None:
    if not batch_records:
        return

    encoded = tokenizer(
        batch_texts,
        add_special_tokens=False,
        truncation=False,
        padding=False,
        return_attention_mask=False,
        return_token_type_ids=False,
    )

    for record, token_ids in zip(batch_records, encoded["input_ids"]):
        content_token_count = len(token_ids)
        if content_token_count == 0:
            metrics.filtered_empty += 1
            continue

        original_token_count = content_token_count + special_tokens
        record["original_token_count"] = original_token_count
        record["original_text_preview"] = record["text"][:280]

        if original_token_count <= max_tokens:
            chunks = [create_single_chunk(record["text"], content_token_count, special_tokens)]
        else:
            metrics.long_documents += 1
            chunks = token_chunk_from_ids(
                token_ids=token_ids,
                tokenizer=tokenizer,
                max_tokens=max_tokens,
                overlap_tokens=overlap_tokens,
                special_tokens=special_tokens,
                start_offset=0,
            )

        emit_document(record, chunks, writer, metrics, sample_state, sample_size, rng)

    batch_records.clear()
    batch_texts.clear()


def segment_corporate_document(
    paragraphs: list[str],
    tokenizer,
    special_tokens: int,
    max_tokens: int,
    overlap_tokens: int,
) -> tuple[list[Chunk], int, bool]:
    cleaned_paragraphs = [paragraph for paragraph in paragraphs if paragraph]
    if not cleaned_paragraphs:
        return [], 0, False

    separator_text = "\n\n"
    separator_ids = tokenizer(separator_text, add_special_tokens=False)["input_ids"]
    separator_length = len(separator_ids)

    encoded = tokenizer(
        cleaned_paragraphs,
        add_special_tokens=False,
        truncation=False,
        padding=False,
        return_attention_mask=False,
        return_token_type_ids=False,
    )
    paragraph_token_ids = encoded["input_ids"]

    chunks: list[Chunk] = []
    current_paragraphs: list[str] = []
    current_start: int | None = None
    current_end: int | None = None
    total_content_tokens = 0
    cursor = 0
    saw_overflow_paragraph = False

    def flush_current() -> None:
        nonlocal current_paragraphs, current_start, current_end
        if not current_paragraphs or current_start is None or current_end is None:
            return
        text = "\n\n".join(current_paragraphs)
        chunks.append(
            Chunk(
                text=text,
                token_count=current_end - current_start + 1 + special_tokens,
                chunk_start_token=current_start,
                chunk_end_token=current_end,
            )
        )
        current_paragraphs = []
        current_start = None
        current_end = None

    for paragraph, paragraph_ids in zip(cleaned_paragraphs, paragraph_token_ids):
        if not paragraph_ids:
            continue

        paragraph_start = cursor
        paragraph_end = paragraph_start + len(paragraph_ids) - 1
        cursor = paragraph_end + 1
        total_content_tokens = cursor

        if len(paragraph_ids) + special_tokens > max_tokens:
            saw_overflow_paragraph = True
            flush_current()
            chunks.extend(
                token_chunk_from_ids(
                    token_ids=paragraph_ids,
                    tokenizer=tokenizer,
                    max_tokens=max_tokens,
                    overlap_tokens=overlap_tokens,
                    special_tokens=special_tokens,
                    start_offset=paragraph_start,
                )
            )
            cursor += separator_length
            total_content_tokens = cursor
            continue

        if not current_paragraphs:
            current_paragraphs = [paragraph]
            current_start = paragraph_start
            current_end = paragraph_end
        else:
            assert current_start is not None and current_end is not None
            projected_end = paragraph_end
            projected_total = projected_end - current_start + 1 + special_tokens
            if projected_total <= max_tokens:
                current_paragraphs.append(paragraph)
                current_end = paragraph_end
            else:
                flush_current()
                current_paragraphs = [paragraph]
                current_start = paragraph_start
                current_end = paragraph_end

        cursor += separator_length
        total_content_tokens = cursor

    flush_current()
    total_content_tokens = max(total_content_tokens - separator_length, 0) if cleaned_paragraphs else 0
    return chunks, total_content_tokens + special_tokens, saw_overflow_paragraph


def process_media_and_academic(
    files,
    tokenizer,
    special_tokens: int,
    max_tokens: int,
    overlap_tokens: int,
    batch_size: int,
    writer: ParquetBufferWriter,
    metrics_by_source: dict[str, SourceMetrics],
    sample_states: dict[str, SampleState],
    sample_size: int,
    rng: random.Random,
) -> None:
    for source in ("media", "academic"):
        metrics = metrics_by_source[source]
        sample_state = sample_states[source]
        batch_records: list[dict[str, Any]] = []
        batch_texts: list[str] = []

        for source_file in files[source]:
            LOGGER.info("Processing %s", source_file.path.name)
            for line_number, raw_row in iter_jsonl(source_file.path):
                metrics.raw_records += 1
                record = build_standard_record(source, raw_row, source_file.path, line_number)

                if record["year"] < 2000:
                    metrics.filtered_year += 1
                    continue
                if not record["text"]:
                    metrics.filtered_empty += 1
                    continue

                batch_records.append(record)
                batch_texts.append(record["text"])

                if len(batch_records) >= batch_size:
                    flush_non_corporate_batch(
                        batch_records=batch_records,
                        batch_texts=batch_texts,
                        tokenizer=tokenizer,
                        special_tokens=special_tokens,
                        max_tokens=max_tokens,
                        overlap_tokens=overlap_tokens,
                        writer=writer,
                        metrics=metrics,
                        sample_state=sample_state,
                        sample_size=sample_size,
                        rng=rng,
                    )

                if metrics.raw_records % 50_000 == 0:
                    LOGGER.info(
                        "[%s] raw=%s kept=%s chunks=%s",
                        source,
                        metrics.raw_records,
                        metrics.kept_documents,
                        metrics.final_chunks,
                    )

        flush_non_corporate_batch(
            batch_records=batch_records,
            batch_texts=batch_texts,
            tokenizer=tokenizer,
            special_tokens=special_tokens,
            max_tokens=max_tokens,
            overlap_tokens=overlap_tokens,
            writer=writer,
            metrics=metrics,
            sample_state=sample_state,
            sample_size=sample_size,
            rng=rng,
        )


def process_corporate(
    files,
    tokenizer,
    special_tokens: int,
    max_tokens: int,
    overlap_tokens: int,
    writer: ParquetBufferWriter,
    metrics_by_source: dict[str, SourceMetrics],
    sample_states: dict[str, SampleState],
    sample_size: int,
    rng: random.Random,
) -> None:
    metrics = metrics_by_source["corporate"]
    sample_state = sample_states["corporate"]

    for source_file in files["corporate"]:
        LOGGER.info("Processing %s", source_file.path.name)
        for line_number, raw_row in iter_jsonl(source_file.path):
            metrics.raw_records += 1
            record = build_standard_record("corporate", raw_row, source_file.path, line_number)

            if record["year"] < 2000:
                metrics.filtered_year += 1
                continue
            if record["sic2"] not in ALLOWED_SIC2:
                metrics.filtered_sic += 1
                continue

            paragraphs = extract_html_paragraphs(record["text"])
            cleaned_text = join_paragraphs(paragraphs)
            if not cleaned_text:
                metrics.filtered_empty += 1
                continue

            chunks, original_token_count, saw_overflow_paragraph = segment_corporate_document(
                paragraphs=paragraphs,
                tokenizer=tokenizer,
                special_tokens=special_tokens,
                max_tokens=max_tokens,
                overlap_tokens=overlap_tokens,
            )
            if not chunks:
                metrics.filtered_empty += 1
                continue

            record["text"] = cleaned_text
            record["original_token_count"] = original_token_count
            record["original_text_preview"] = cleaned_text[:280]

            if original_token_count > max_tokens:
                metrics.long_documents += 1
            if saw_overflow_paragraph:
                metrics.paragraph_overflow_documents += 1

            emit_document(record, chunks, writer, metrics, sample_state, sample_size, rng)

            if metrics.raw_records % 500 == 0:
                LOGGER.info(
                    "[corporate] raw=%s kept=%s chunks=%s",
                    metrics.raw_records,
                    metrics.kept_documents,
                    metrics.final_chunks,
                )


def autosize_worksheet(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 80)


def write_validation_workbook(output_path: Path, sample_states: dict[str, SampleState]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheet_mapping = {
        "media": "guardian_chunks",
        "academic": "papers_chunks",
        "corporate": "10k_chunks",
    }
    headers = [
        "source",
        "doc_id",
        "chunk_id",
        "year",
        "company",
        "sic",
        "sic2",
        "industry",
        "item",
        "token_count",
        "chunk_start_token",
        "chunk_end_token",
        "original_text_preview",
        "chunk_text",
        "validation_status",
        "validation_ok",
        "reviewer_notes",
    ]

    for source in ("media", "academic", "corporate"):
        worksheet = workbook.create_sheet(sheet_mapping[source])
        worksheet.append(headers)
        rows = sample_states[source].reservoir
        for row in sorted(rows, key=lambda item: (item["doc_id"], item["chunk_id"])):
            worksheet.append([row.get(header, "") for header in headers])

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top")
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        autosize_worksheet(worksheet)

    workbook.save(output_path)
    LOGGER.info("Wrote %s", output_path)


def format_ratio(numerator: int, denominator: int) -> str:
    if denominator == 0:
        return "0.00"
    return f"{numerator / denominator:.2f}"


def write_chunking_report(
    output_path: Path,
    tokenizer_origin: str,
    max_tokens: int,
    overlap_tokens: int,
    metrics_by_source: dict[str, SourceMetrics],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    rows = []
    for source in ("media", "academic", "corporate"):
        metrics = metrics_by_source[source]
        rows.append(
            "| "
            + " | ".join(
                [
                    source,
                    str(metrics.raw_records),
                    str(metrics.filtered_year),
                    str(metrics.filtered_sic),
                    str(metrics.filtered_empty),
                    str(metrics.kept_documents),
                    str(metrics.chunked_documents),
                    str(metrics.final_chunks),
                    format_ratio(metrics.final_chunks, metrics.kept_documents),
                ]
            )
            + " |"
        )

    observations = [
        "- `media`: Guardian rows were treated as paragraph-level documents, consistent with the stage 0 inspection. Only documents above the token threshold were split with overlap.",
        "- `academic`: abstracts were kept whole when possible and split only when the BGE tokenizer produced more than 350 tokens.",
        "- `corporate`: raw HTML in `content` was converted into cleaned paragraph blocks first. Paragraphs were packed into chunks up to the token limit, and only paragraphs that still exceeded the limit were token-split with overlap.",
        f"- Global scope filter applied: `year >= 2000`.",
        f"- Corporate sector filter applied after computing `sic2 = int(sic) // 100`, retaining only {', '.join(sorted(ALLOWED_SIC2))}.",
    ]

    corporate_metrics = metrics_by_source["corporate"]
    observations.append(
        f"- Corporate documents with at least one over-limit paragraph after HTML cleaning: {corporate_metrics.paragraph_overflow_documents}."
    )

    content = [
        "# Chunking Report",
        "",
        f"- Generated at: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S %Z')}",
        f"- Tokenizer: `{TOKENIZER_REPO_ID}`",
        f"- Tokenizer loaded from: `{tokenizer_origin}`",
        f"- Chunking parameters: `max_tokens={max_tokens}`, `overlap_tokens={overlap_tokens}`",
        "",
        "## Source Summary",
        "",
        "| source | raw_records | filtered_year | filtered_sic | filtered_empty | kept_documents | chunked_documents | final_chunks | expansion_factor |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- |",
        *rows,
        "",
        "## Observations",
        "",
        *observations,
        "",
    ]

    output_path.write_text("\n".join(content), encoding="utf-8")
    LOGGER.info("Wrote %s", output_path)


def main() -> int:
    args = parse_args()
    logging.basicConfig(
        level=getattr(logging, args.log_level.upper(), logging.INFO),
        format="%(asctime)s %(levelname)s %(message)s",
    )

    args.project_root.mkdir(parents=True, exist_ok=True)
    rng = random.Random(args.seed)

    tokenizer, tokenizer_origin = load_tokenizer(args.tokenizer_path)
    files = locate_input_files(args.project_root)
    special_tokens = tokenizer.num_special_tokens_to_add(pair=False)
    output_corpus = args.project_root / OUTPUT_CORPUS_REL
    chunking_report = args.project_root / CHUNKING_REPORT_REL
    validation_workbook = args.project_root / VALIDATION_WORKBOOK_REL

    writer = ParquetBufferWriter(output_corpus, FINAL_SCHEMA, args.parquet_batch_size)
    metrics_by_source = {
        "media": SourceMetrics(),
        "academic": SourceMetrics(),
        "corporate": SourceMetrics(),
    }
    sample_states = {
        "media": SampleState(),
        "academic": SampleState(),
        "corporate": SampleState(),
    }

    process_media_and_academic(
        files=files,
        tokenizer=tokenizer,
        special_tokens=special_tokens,
        max_tokens=args.max_tokens,
        overlap_tokens=args.overlap_tokens,
        batch_size=args.batch_size,
        writer=writer,
        metrics_by_source=metrics_by_source,
        sample_states=sample_states,
        sample_size=args.validation_sample_size,
        rng=rng,
    )
    process_corporate(
        files=files,
        tokenizer=tokenizer,
        special_tokens=special_tokens,
        max_tokens=args.max_tokens,
        overlap_tokens=args.overlap_tokens,
        writer=writer,
        metrics_by_source=metrics_by_source,
        sample_states=sample_states,
        sample_size=args.validation_sample_size,
        rng=rng,
    )

    writer.close()
    write_chunking_report(
        output_path=chunking_report,
        tokenizer_origin=tokenizer_origin,
        max_tokens=args.max_tokens,
        overlap_tokens=args.overlap_tokens,
        metrics_by_source=metrics_by_source,
    )
    write_validation_workbook(validation_workbook, sample_states)
    LOGGER.info("Corpus preparation completed successfully.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
